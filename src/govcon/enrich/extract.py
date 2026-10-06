"""Document text extraction for PDF, DOCX, XLSX, and plain text.

Text comes from each file's own text layer. PDF pages without one are read
by OCR when an ``OcrConfig`` is passed (ADR-065); every result also carries
its pages, so citations can name a page. Extraction status and errors are
tracked per file so downstream consumers know what succeeded and what needs
manual review.
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, TypedDict

if TYPE_CHECKING:
    from govcon.enrich.ocr import OcrConfig

logger = logging.getLogger("govcon.enrich.extract")


@dataclass(frozen=True)
class ExtractionLimits:
    """Bounds that keep a hostile or malformed file from exhausting the host."""

    max_pdf_pages: int = 1500
    max_text_chars: int = 5_000_000
    max_zip_members: int = 5_000
    max_zip_uncompressed_bytes: int = 300 * 1024 * 1024
    max_zip_compression_ratio: float = 200.0
    max_xlsx_sheets: int = 50
    max_xlsx_cells: int = 500_000


DEFAULT_LIMITS = ExtractionLimits()


class ExtractionLimitExceeded(ValueError):
    """The file is rejected before parsing (archive bomb or similar)."""


def check_zip_container(data: bytes, limits: ExtractionLimits = DEFAULT_LIMITS) -> None:
    """Reject DOCX/XLSX (ZIP) containers that would expand unreasonably."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionLimitExceeded(f"not a valid ZIP container: {exc}") from exc
    with archive:
        members = archive.infolist()
        if len(members) > limits.max_zip_members:
            raise ExtractionLimitExceeded(f"{len(members)} archive members exceed the {limits.max_zip_members} limit")
        total = 0
        for member in members:
            total += member.file_size
            if member.compress_size > 0 and member.file_size / member.compress_size > limits.max_zip_compression_ratio:
                raise ExtractionLimitExceeded(f"member {member.filename!r} has a suspicious compression ratio")
            if member.compress_size == 0 and member.file_size > 0:
                raise ExtractionLimitExceeded(f"member {member.filename!r} has a suspicious compression ratio")
        if total > limits.max_zip_uncompressed_bytes:
            raise ExtractionLimitExceeded(f"archive expands to {total} bytes, over the limit")


def _cap_text(result: "ExtractionResult", limits: ExtractionLimits) -> "ExtractionResult":
    if result.text is not None and len(result.text) > limits.max_text_chars:
        note = f"text truncated at {limits.max_text_chars} characters"
        return ExtractionResult(
            result.text[: limits.max_text_chars],
            "partial",
            f"{result.error}; {note}" if result.error else note,
            page_count=result.page_count,
            pages=result.pages,
            ocr_pages=result.ocr_pages,
            ocr_failed_pages=result.ocr_failed_pages,
        )
    return result


@dataclass(frozen=True)
class PageText:
    """One page (PDF) or part (sheet, whole document) of a file's text."""

    page_no: int
    text: str
    source: str = "native"  # native | ocr | none
    confidence: float | None = None
    label: str | None = None


class _PageCoverage(TypedDict):
    page_count: int
    pages: list[PageText]
    ocr_pages: list[int]
    ocr_failed_pages: list[dict]


class ExtractionResult:
    __slots__ = ("text", "status", "error", "page_count", "pages", "ocr_pages", "ocr_failed_pages")

    def __init__(
        self,
        text: str | None,
        status: str,
        error: str | None = None,
        page_count: int | None = None,
        pages: list[PageText] | None = None,
        ocr_pages: list[int] | None = None,
        ocr_failed_pages: list[dict] | None = None,
    ):
        self.text = text
        self.status = status
        self.error = error
        self.page_count = page_count
        self.pages = pages
        self.ocr_pages = ocr_pages or []
        self.ocr_failed_pages = ocr_failed_pages or []


def extract_text(
    data: bytes,
    mime_type: str,
    filename: str | None = None,
    limits: ExtractionLimits = DEFAULT_LIMITS,
    ocr: "OcrConfig | None" = None,
) -> ExtractionResult:
    """Extract text from file bytes based on MIME type.

    Returns an ``ExtractionResult`` with status ``success``, ``partial``,
    ``unsupported``, or ``error``, and the file's pages. ``limits`` bound
    pages, archive expansion, workbook size, and returned text; truncation
    yields ``partial``. With ``ocr``, PDF pages lacking a text layer are OCR'd.
    """
    mime = (mime_type or "").lower().split(";", 1)[0].strip()
    fname = (filename or "").lower()

    if mime == "application/pdf" or fname.endswith(".pdf"):
        return _cap_text(_extract_pdf(data, limits, ocr), limits)
    if mime in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ) or fname.endswith(".docx"):
        return _cap_text(_whole_document(_extract_docx(data, limits)), limits)
    if mime in (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    ) or fname.endswith(".xlsx"):
        return _cap_text(_sheet_pages(_extract_xlsx(data, limits)), limits)
    if mime.startswith("text/") or fname.endswith(".txt") or fname.endswith(".csv"):
        return _cap_text(_whole_document(_extract_plain(data)), limits)

    return ExtractionResult(None, "unsupported", f"unsupported mime type: {mime}")


def _whole_document(result: ExtractionResult) -> ExtractionResult:
    """A file without pages is one part, labelled ``document``."""
    if result.text is not None:
        result.pages = [PageText(1, result.text, "native", label="document")]
    return result


def _sheet_pages(result: ExtractionResult) -> ExtractionResult:
    """One part per worksheet, labelled with the sheet name."""
    if result.text is None:
        return result
    pages: list[PageText] = []
    for number, block in enumerate(result.text.split("\n\n[Sheet: "), start=1):
        body = block if number == 1 else "[Sheet: " + block
        name = body.removeprefix("[Sheet: ").split("]", 1)[0]
        pages.append(PageText(number, body, "native", label=f"sheet {name}"))
    result.pages = pages
    return result


def _extract_pdf(data: bytes, limits: ExtractionLimits = DEFAULT_LIMITS,
                 ocr: "OcrConfig | None" = None) -> ExtractionResult:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        total_pages = len(reader.pages)
        native: list[str] = []
        for index, page in enumerate(reader.pages):
            if index >= limits.max_pdf_pages:
                break
            native.append(page.extract_text() or "")
    except Exception as exc:
        logger.warning("PDF extraction failed: %s", exc)
        return ExtractionResult(None, "error", str(exc))

    pages = [PageText(n, text, "native" if text.strip() else "none") for n, text in enumerate(native, start=1)]
    ocr_pages: list[int] = []
    ocr_failed: list[dict] = []
    if ocr is not None and ocr.enabled:
        from govcon.enrich.ocr import ocr_pdf_pages

        candidates = [p.page_no for p in pages if len(p.text.strip()) < ocr.min_native_chars]
        read, failed = ocr_pdf_pages(data, candidates, ocr)
        for page_no, ocr_page in read.items():
            pages[page_no - 1] = PageText(page_no, ocr_page.text, "ocr", ocr_page.confidence)
            ocr_pages.append(page_no)
        # A page that keeps some native text is readable; only blank pages count as failed.
        ocr_failed = [{"page": page_no, "reason": reason} for page_no, reason in sorted(failed.items())
                      if not pages[page_no - 1].text.strip()]

    full_text = "\n\n".join(p.text for p in pages)
    extras: _PageCoverage = {"page_count": total_pages, "pages": pages, "ocr_pages": sorted(ocr_pages), "ocr_failed_pages": ocr_failed}
    if not full_text.strip():
        reason = f"; OCR: {ocr_failed[0]['reason']}" if ocr_failed else ""
        return ExtractionResult(None, "partial", f"PDF contained no extractable text (may need OCR){reason}", **extras)
    if total_pages > limits.max_pdf_pages:
        return ExtractionResult(
            full_text, "partial", f"PDF truncated at {limits.max_pdf_pages} of {total_pages} pages", **extras
        )
    return ExtractionResult(full_text, "success", **extras)


def extract_pdf_pages(data: bytes, limits: ExtractionLimits = DEFAULT_LIMITS) -> list[str] | None:
    """Return per-page text for a PDF (at most ``max_pdf_pages``), or ``None`` when unreadable."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [page.extract_text() or "" for _, page in zip(range(limits.max_pdf_pages), reader.pages)]
    except Exception as exc:
        logger.warning("PDF page extraction failed: %s", exc)
        return None


def _extract_docx(data: bytes, limits: ExtractionLimits = DEFAULT_LIMITS) -> ExtractionResult:
    try:
        check_zip_container(data, limits)
    except ExtractionLimitExceeded as exc:
        return ExtractionResult(None, "error", f"rejected before parsing: {exc}")
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Requirements are often embedded in tables (packaging, CLIN, marking);
        # paragraph-only extraction silently drops them.
        tables: list[str] = []
        for index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                tables.append(f"[Table {index}]\n" + "\n".join(rows))
        if not paragraphs and not tables:
            return ExtractionResult(None, "partial", "DOCX contained no text paragraphs or tables")
        return ExtractionResult("\n\n".join(paragraphs + tables), "success")
    except Exception as exc:
        logger.warning("DOCX extraction failed: %s", exc)
        return ExtractionResult(None, "error", str(exc))


def _extract_xlsx(data: bytes, limits: ExtractionLimits = DEFAULT_LIMITS) -> ExtractionResult:
    try:
        check_zip_container(data, limits)
    except ExtractionLimitExceeded as exc:
        return ExtractionResult(None, "error", f"rejected before parsing: {exc}")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        notes: list[str] = []
        cells_seen = 0
        sheet_names = list(wb.sheetnames)
        if len(sheet_names) > limits.max_xlsx_sheets:
            notes.append(f"only the first {limits.max_xlsx_sheets} of {len(sheet_names)} sheets were read")
            sheet_names = sheet_names[: limits.max_xlsx_sheets]
        for sheet in sheet_names:
            ws = wb[sheet]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells_seen += len(row)
                if cells_seen > limits.max_xlsx_cells:
                    notes.append(f"stopped after {limits.max_xlsx_cells} cells")
                    break
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                parts.append(f"[Sheet: {sheet}]\n" + "\n".join(rows))
            if cells_seen > limits.max_xlsx_cells:
                break
        wb.close()
        if not parts:
            return ExtractionResult(None, "partial", "XLSX contained no data")
        if notes:
            return ExtractionResult("\n\n".join(parts), "partial", "; ".join(notes))
        return ExtractionResult("\n\n".join(parts), "success")
    except Exception as exc:
        logger.warning("XLSX extraction failed: %s", exc)
        return ExtractionResult(None, "error", str(exc))


def _extract_plain(data: bytes) -> ExtractionResult:
    try:
        for encoding in ("utf-8", "latin-1"):
            try:
                text = data.decode(encoding)
                return ExtractionResult(text, "success")
            except UnicodeDecodeError:
                continue
        return ExtractionResult(None, "error", "unable to decode text file")
    except Exception as exc:
        return ExtractionResult(None, "error", str(exc))


def guess_mime_type(filename: str) -> str:
    """Return a MIME type based on file extension."""
    ext = Path(filename).suffix.lower()
    return {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".doc": "application/msword",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xls": "application/vnd.ms-excel",
        ".txt": "text/plain",
        ".csv": "text/csv",
        ".html": "text/html",
        ".htm": "text/html",
    }.get(ext, "application/octet-stream")
