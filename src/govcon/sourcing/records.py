"""Supplier, product, catalog, quote and RFQ records (roadmap gap 6, ADR-071).

Every record keeps its provenance (who entered it, or which file it came
from) and quotes keep their validity date. Nothing here contacts a supplier:
an RFQ is a draft for a person to send.
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from govcon.audit import record_audit
from govcon.collaboration.users import require_permission
from govcon.matching.pricing import canonical_nsn
from govcon.models import (
    AISharingAuthorization,
    CatalogImport,
    Opportunity,
    Product,
    RfqDraft,
    Supplier,
    SupplierProduct,
    SupplierQuote,
    SupplierQuoteLine,
    User,
)

QUOTE_SCOPE = "supplier_quotes"
SOURCING_REQUIREMENT_TYPES = frozenset(
    {"technical", "delivery", "country_of_origin", "packaging", "quality", "certification", "item", "marking"}
)


class SourcingError(ValueError):
    pass


def _decimal(value: Any, what: str) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).replace("$", "").replace(",", "").strip())
    except InvalidOperation as exc:
        raise SourcingError(f"{what} is not a number: {value!r}") from exc
    if not number.is_finite():
        raise SourcingError(f"{what} must be finite: {value!r}")
    if number < 0:
        raise SourcingError(f"{what} cannot be negative: {value!r}")
    return number


def _date(value: Any, what: str) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as exc:
        raise SourcingError(f"{what} must be a date (YYYY-MM-DD): {value!r}") from exc


# ── suppliers and catalogs ───────────────────────────────────────────────────

def get_or_create_supplier(session: Session, *, name: str, actor: User | None, provenance: str,
                           **fields: Any) -> tuple[Supplier, bool]:
    clean = (name or "").strip()
    if not clean:
        raise SourcingError("a supplier needs a name")
    existing = session.scalar(select(Supplier).where(func.lower(Supplier.name) == clean.lower()))
    if existing is not None:
        return existing, False
    supplier = Supplier(name=clean, provenance=provenance, created_by_user_id=actor.id if actor else None,
                        **{k: (v.strip() if isinstance(v, str) else v) or None for k, v in fields.items()})
    session.add(supplier)
    session.flush()
    record_audit(session, action_type="supplier_created", user_id=actor.id if actor else None,
                 entity_type="suppliers", entity_id=supplier.id, new_value={"name": clean, "provenance": provenance})
    return supplier, True


@dataclass
class CatalogResult:
    import_id: int
    rows_total: int
    rows_imported: int
    errors: list[dict[str, Any]] = field(default_factory=list)


def import_catalog_csv(session: Session, *, supplier_id: int, data: bytes, filename: str | None,
                       actor: User) -> CatalogResult:
    """Import a supplier catalog CSV.

    Columns: ``part_number`` (required), ``manufacturer``, ``nsn``,
    ``description``, ``unit``, ``list_price``, ``valid_until``. Invalid rows are
    reported with their line number and skipped; valid rows are upserted.
    """
    require_permission(actor, "review")
    supplier = session.get(Supplier, supplier_id)
    if supplier is None:
        raise SourcingError("supplier not found")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise SourcingError("the catalog must be a UTF-8 CSV file") from exc
    reader = csv.DictReader(io.StringIO(text))
    headers = {h.strip().lower() for h in reader.fieldnames or []}
    if "part_number" not in headers:
        raise SourcingError("the catalog CSV needs a part_number column")
    run = CatalogImport(supplier_id=supplier.id, filename=filename, sha256=hashlib.sha256(data).hexdigest(),
                        imported_by_user_id=actor.id)
    session.add(run)
    session.flush()
    errors: list[dict[str, Any]] = []
    total = imported = 0
    for line, raw in enumerate(reader, start=2):
        total += 1
        try:
            if None in raw:
                raise SourcingError("row has more fields than the catalog headers")
            row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw.items()}
            part = row.get("part_number")
            if not part:
                raise SourcingError("part_number is empty")
            price = _decimal(row.get("list_price"), "list_price")
            valid_until = _date(row.get("valid_until"), "valid_until")
            manufacturer = row.get("manufacturer") or None
            product = session.scalar(select(Product).where(
                func.coalesce(func.lower(Product.manufacturer), "") == (manufacturer or "").lower(),
                func.lower(Product.part_number) == part.lower()))
            if product is None:
                product = Product(part_number=part, manufacturer=manufacturer)
                session.add(product)
            product.nsn = canonical_nsn(row.get("nsn")) or product.nsn
            product.description = row.get("description") or product.description
            product.unit = row.get("unit") or product.unit
            session.flush()
            link = session.scalar(select(SupplierProduct).where(
                SupplierProduct.supplier_id == supplier.id, SupplierProduct.product_id == product.id))
            if link is None:
                link = SupplierProduct(supplier_id=supplier.id, product_id=product.id)
                session.add(link)
            link.list_price, link.valid_until, link.catalog_import_id = price, valid_until, run.id
            session.flush()
            imported += 1
        except SourcingError as exc:
            errors.append({"line": line, "error": str(exc)})
    run.rows_total, run.rows_imported, run.errors = total, imported, errors
    record_audit(session, action_type="catalog_imported", user_id=actor.id, entity_type="catalog_imports",
                 entity_id=run.id, new_value={"supplier_id": supplier.id, "rows": total, "imported": imported,
                                              "errors": len(errors)})
    session.flush()
    return CatalogResult(run.id, total, imported, errors)


# ── quotes ───────────────────────────────────────────────────────────────────

QUOTE_COLUMNS = ("description", "part_number", "nsn", "quantity", "unit", "unit_price", "extended_price", "lead_time_days")


def parse_quote_table(data: bytes, filename: str) -> list[dict[str, Any]]:
    """Quote lines from a CSV or XLSX table with the ``QUOTE_COLUMNS`` headers. No AI."""
    name = filename.lower()
    if name.endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from zipfile import BadZipFile

        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = workbook.active
            if sheet is None:
                raise SourcingError("the quote workbook has no active worksheet")
            values = list(sheet.iter_rows(values_only=True))
        except (BadZipFile, InvalidFileException, KeyError, ValueError, SyntaxError, OSError) as exc:
            raise SourcingError("the quote workbook is invalid or corrupt") from exc
        finally:
            if workbook is not None:
                workbook.close()
        if not values:
            return []
        headers = [str(h or "").strip().lower() for h in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:] if any(v not in (None, "") for v in row)]
    else:
        raise SourcingError("only CSV or XLSX quote tables are parsed without AI")
    lines = []
    for number, raw in enumerate(rows, start=2):
        row = {str(k or "").strip().lower(): v for k, v in raw.items()}
        if not any(row.get(c) not in (None, "") for c in ("description", "part_number", "nsn", "unit_price")):
            continue
        try:
            lead = row.get("lead_time_days")
            lines.append({
                "description": (str(row.get("description") or "").strip() or None),
                "part_number": (str(row.get("part_number") or "").strip() or None),
                "nsn": canonical_nsn(str(row.get("nsn") or "")) or None,
                "quantity": _decimal(row.get("quantity"), f"line {number} quantity"),
                "unit": (str(row.get("unit") or "").strip() or None),
                "unit_price": _decimal(row.get("unit_price"), f"line {number} unit_price"),
                "extended_price": _decimal(row.get("extended_price"), f"line {number} extended_price"),
                "lead_time_days": int(lead) if lead not in (None, "") else None,
            })
        except (TypeError, ValueError) as exc:
            raise SourcingError(f"line {number}: {exc}") from exc
    return lines


def record_quote(session: Session, *, opportunity_id: int, supplier_id: int, actor: User | None, method: str,
                 lines: list[dict[str, Any]] | None = None, total_price: Any = None, valid_until: Any = None,
                 source: dict[str, Any] | None = None, notes: str | None = None) -> SupplierQuote:
    """Store a quote with its lines. Totals are computed from lines, never guessed."""
    from govcon.workflow.invalidation import lock_opportunity

    if actor is not None:
        require_permission(actor, "review")
    if lock_opportunity(session, opportunity_id) is None or session.get(Supplier, supplier_id) is None:
        raise SourcingError("opportunity or supplier not found")
    lines = lines or []
    total = _decimal(total_price, "total_price")
    if total is None and lines:
        extended = []
        for line in lines:
            value = line.get("extended_price")
            if value is None and line.get("unit_price") is not None and line.get("quantity") is not None:
                value = Decimal(str(line["unit_price"])) * Decimal(str(line["quantity"]))
            extended.append(value)
        total = sum(extended, Decimal(0)) if all(v is not None for v in extended) else None
    validity = _date(valid_until, "valid_until")
    if source and source.get("source_sha256"):
        existing = session.scalar(select(SupplierQuote).where(
            SupplierQuote.opportunity_id == opportunity_id, SupplierQuote.supplier_id == supplier_id,
            SupplierQuote.source_sha256 == source["source_sha256"], SupplierQuote.valid_until == validity,
            SupplierQuote.status == "active").order_by(SupplierQuote.id).limit(1))
        if existing is not None:
            return existing
    quote = SupplierQuote(
        opportunity_id=opportunity_id, supplier_id=supplier_id, extraction_method=method, total_price=total,
        valid_until=validity, entered_by_user_id=actor.id if actor else None,
        notes=notes, **(source or {}),
    )
    session.add(quote)
    session.flush()
    for line in lines:
        session.add(SupplierQuoteLine(quote_id=quote.id, **{k: line.get(k) for k in QUOTE_COLUMNS}))
    record_audit(session, action_type="supplier_quote_recorded", user_id=actor.id if actor else None,
                 opportunity_id=opportunity_id, entity_type="supplier_quotes", entity_id=quote.id,
                 new_value={"supplier_id": supplier_id, "method": method, "lines": len(lines),
                            "total_price": str(total) if total is not None else None,
                            "valid_until": quote.valid_until.isoformat() if quote.valid_until else None})
    session.flush()
    return quote


def store_quote_file(data: bytes, filename: str, opportunity_id: int, settings) -> dict[str, Any]:
    """Keep the quote document's bytes in the attachment store, outside the solicitation files."""
    from govcon.enrich.attachments import store_bytes

    sha = hashlib.sha256(data).hexdigest()
    key = str(store_bytes(data, sha, opportunity_id, f"quote_{filename}", settings))
    return {"source_filename": filename, "source_sha256": sha, "source_key": key}


def is_current(quote: SupplierQuote, today: date | None = None) -> bool:
    today = today or datetime.now(UTC).date()
    return quote.status == "active" and (quote.valid_until is None or quote.valid_until >= today)


def current_quotes(session: Session, opportunity_id: int) -> list[SupplierQuote]:
    quotes = session.scalars(select(SupplierQuote).where(SupplierQuote.opportunity_id == opportunity_id)
                             .order_by(SupplierQuote.id)).all()
    return [q for q in quotes if is_current(q)]


def quote_records(session: Session, opportunity_id: int) -> list[dict[str, Any]]:
    """Current quotes with lines, as source records for supplier and pricing analysis."""
    records = []
    for quote in current_quotes(session, opportunity_id):
        supplier = session.get(Supplier, quote.supplier_id)
        lines = session.scalars(select(SupplierQuoteLine).where(SupplierQuoteLine.quote_id == quote.id)).all()
        records.append({
            "quote_id": quote.id,
            "supplier": supplier.name if supplier else None,
            "total_price": float(quote.total_price) if quote.total_price is not None else None,
            "valid_until": quote.valid_until.isoformat() if quote.valid_until else "UNKNOWN",
            "received_at": quote.received_at.isoformat() if quote.received_at else None,
            "source": f"supplier quote ({quote.extraction_method}"
                      + (f", file {quote.source_filename}" if quote.source_filename else "") + ")",
            "lines": [{k: (float(v) if isinstance(v, Decimal) else v) for k, v in
                       ((c, getattr(line, c)) for c in QUOTE_COLUMNS)} for line in lines],
        })
    return records


def lowest_current_total(session: Session, opportunity_id: int) -> SupplierQuote | None:
    opportunity = session.get(Opportunity, opportunity_id)
    priced = []
    for quote in current_quotes(session, opportunity_id):
        if quote.total_price is None or opportunity is None:
            continue
        lines = list(session.scalars(select(SupplierQuoteLine).where(SupplierQuoteLine.quote_id == quote.id)))
        if lines:
            if opportunity.quantity is None or any(line.quantity is None for line in lines):
                continue
            if sum((line.quantity for line in lines), Decimal(0)) != opportunity.quantity:
                continue
            if not opportunity.unit or any((line.unit or "").strip().upper() != opportunity.unit.strip().upper()
                                           for line in lines):
                continue
            if opportunity.nsn and any(line.nsn and canonical_nsn(line.nsn) != canonical_nsn(opportunity.nsn)
                                       for line in lines):
                continue
        elif quote.extraction_method != "manual":
            continue
        priced.append(quote)
    return min(priced, key=lambda q: q.total_price, default=None)


def sourcing_revision(session: Session, opportunity_id: int) -> str:
    """Hash of the current quotes; a change makes queued supplier/pricing analyses stale."""
    basis = [(q.id, str(q.total_price), q.valid_until.isoformat() if q.valid_until else None, q.status)
             for q in current_quotes(session, opportunity_id)]
    return hashlib.sha256(repr(basis).encode()).hexdigest()


# ── RFQ drafts ───────────────────────────────────────────────────────────────

def draft_rfq(session: Session, *, opportunity_id: int, supplier_id: int | None, actor: User) -> RfqDraft:
    """Draft a request for quote from the opportunity's own facts. It is never sent."""
    from govcon.compliance.matrix import active_requirements

    require_permission(actor, "review")
    opp = session.get(Opportunity, opportunity_id)
    if opp is None:
        raise SourcingError("opportunity not found")
    supplier = session.get(Supplier, supplier_id) if supplier_id else None
    needs = [r for r in active_requirements(session, opportunity_id)
             if (r.requirement_type or "") in SOURCING_REQUIREMENT_TYPES]
    lines = [
        f"Dear {supplier.contact_name or supplier.name}," if supplier else "Hello,",
        "",
        f"We are preparing a quote for {opp.agency_path or 'a federal buyer'}"
        f"{' solicitation ' + opp.solicitation_number if opp.solicitation_number else ''}: {opp.title}.",
        "Please quote the following:",
        f"- NSN: {opp.nsn or 'not stated'}",
        f"- Quantity: {opp.quantity if opp.quantity is not None else 'not stated'} {opp.unit or ''}".rstrip(),
    ]
    if needs:
        lines.append("Requirements from the solicitation:")
        lines += [f"- {r.requirement_text}" for r in needs[:25]]
    deadline = opp.response_deadline - timedelta(days=2) if opp.response_deadline else None
    lines += [
        "",
        "Please include unit price, lead time, country of origin, and how long the quote is valid.",
        f"We need your quote by {deadline.strftime('%Y-%m-%d')}." if deadline else "Please reply at your earliest convenience.",
        "",
        "Thank you.",
    ]
    draft = RfqDraft(opportunity_id=opportunity_id, supplier_id=supplier.id if supplier else None,
                     subject=f"RFQ: {opp.solicitation_number or opp.title}"[:200], body="\n".join(lines),
                     created_by_user_id=actor.id)
    session.add(draft)
    session.flush()
    record_audit(session, action_type="rfq_drafted", user_id=actor.id, opportunity_id=opportunity_id,
                 entity_type="rfq_drafts", entity_id=draft.id, new_value={"supplier_id": draft.supplier_id})
    return draft


# ── explicit AI sharing authorizations ───────────────────────────────────────

def grant_authorization(session: Session, *, provider: str, days: int, reason: str, actor: User,
                        scope: str = QUOTE_SCOPE) -> AISharingAuthorization:
    """Owner-only: allow ``scope`` data to go to one AI provider until it expires."""
    require_permission(actor, "manage_users")
    if not 1 <= days <= 365:
        raise SourcingError("an authorization lasts between 1 and 365 days")
    if len((reason or "").strip()) < 10:
        raise SourcingError("give a reason of at least 10 characters")
    row = AISharingAuthorization(scope=scope, provider=provider.strip().lower(), reason=reason.strip(),
                                 granted_by_user_id=actor.id,
                                 expires_at=datetime.now(UTC) + timedelta(days=days))
    session.add(row)
    session.flush()
    record_audit(session, action_type="ai_sharing_authorized", user_id=actor.id,
                 entity_type="ai_sharing_authorizations", entity_id=row.id,
                 new_value={"scope": scope, "provider": row.provider, "expires_at": row.expires_at.isoformat(),
                            "reason": row.reason})
    return row


def revoke_authorization(session: Session, *, authorization_id: int, actor: User) -> AISharingAuthorization:
    require_permission(actor, "manage_users")
    row = session.get(AISharingAuthorization, authorization_id, with_for_update=True)
    if row is None:
        raise SourcingError("authorization not found")
    if row.revoked_at is None:
        row.revoked_at = datetime.now(UTC)
        row.revoked_by_user_id = actor.id
        record_audit(session, action_type="ai_sharing_revoked", user_id=actor.id,
                     entity_type="ai_sharing_authorizations", entity_id=row.id, new_value={"provider": row.provider})
    session.flush()
    return row


def active_authorization(session: Session, *, provider: str, scope: str = QUOTE_SCOPE,
                         now: datetime | None = None) -> AISharingAuthorization | None:
    now = now or datetime.now(UTC)
    return session.scalar(
        select(AISharingAuthorization).where(
            AISharingAuthorization.scope == scope, AISharingAuthorization.provider == provider.strip().lower(),
            AISharingAuthorization.revoked_at.is_(None), AISharingAuthorization.expires_at > now,
        ).order_by(AISharingAuthorization.expires_at.desc()).limit(1)
    )
