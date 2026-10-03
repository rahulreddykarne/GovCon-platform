"""Release acceptance through service gates, with real DOCX/XLSX artifacts.

The function is also executed against the installed wheel, outside the checkout.
Every approved/ready/submitted state comes from an authorized workflow service.
"""
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from pathlib import Path
from threading import Barrier
from uuid import uuid4

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session


def prepared_lifecycle(engine, artifact_dir):
    from docx import Document
    from openpyxl import Workbook
    from govcon.collaboration.users import invite_user
    from govcon.collaboration.assignments import assign_reviewer
    from govcon.collaboration.comments import add_comment
    from govcon.collaboration.review_sessions import complete_assignment, finalize_approval, set_review_policy
    from govcon.compliance.inventory import build_document_inventory
    from govcon.compliance.matrix import override_requirement
    from govcon.compliance.proposal_coverage import check_proposal_coverage
    from govcon.compliance.deterministic import PackageFile, SubmissionPackage
    from govcon.config import Settings
    from govcon.enrich.attachments import process_local_file
    from govcon.models import Opportunity, Proposal, Requirement, Submission
    from govcon.proposals.service import generate_proposal
    from govcon.proposals.versions import create_proposal_version
    from govcon.proposals.export import export_proposal_docx
    from govcon.security.classification import DataClassification
    from govcon.submissions.service import generate_submission_package
    from govcon.submissions.manifest import assemble_package
    configured = Settings(_env_file=None, decision_primary_provider="rules", decision_fallback_provider="rules",
                          deepseek_api_key=None, anthropic_api_key=None, openai_api_key=None, jev_api_key=None,
                          ai_external_allowed_for_proprietary=False, ai_external_allowed_for_fci=False, ai_external_allowed_for_cui=False)
    artifact_dir = Path(artifact_dir)
    artifact_dir.mkdir(parents=True, exist_ok=True)
    text = "Submit proposal.docx and pricing.xlsx by email to release@example.test. The deadline timezone is UTC."
    source_path = artifact_dir / "solicitation.docx"
    source = Document()
    source.add_paragraph(text)
    source.save(source_path)
    from govcon.db import settings_scope
    with settings_scope(configured), Session(engine) as db:
        opp = Opportunity(source="sam", source_id="release-"+uuid4().hex, title="Synthetic release lifecycle", status="open",
                          raw={}, links={}, response_deadline=datetime.now(UTC)+timedelta(days=30))
        db.add(opp)
        db.flush()
        actor = invite_user(db, email=uuid4().hex+"@release.test", display_name="Release approver", password="synthetic strong password", role="approver")
        reviewers = [invite_user(db, email=uuid4().hex+"@release.test", display_name=f"Reviewer {i}", password="synthetic strong password", role="reviewer") for i in range(2)]
        file = process_local_file(db, opp, source_path, classification=DataClassification.PUBLIC, source_origin="synthetic release fixture")
        requirement = Requirement(opportunity_id=opp.id, source_file_id=file.id, requirement_text=text, source_quote=text,
            requirement_type="submission", mandatory=True, severity="high", status="unreviewed", blocks_submission=True,
            key_values={"submission_method":"email", "recipient_email":"release@example.test", "deadline_timezone":"UTC",
                        "required_files":["proposal.docx","pricing.xlsx"], "allowed_file_types":["DOCX","XLSX"]})
        db.add(requirement)
        db.flush()
        inventory, _ = build_document_inventory(db, opp.id)
        assert inventory.complete
        # Initial requirement is unresolved; a documented human verification
        # via the authorized service establishes it, rather than a state write.
        override_requirement(db, requirement_id=requirement.id, status="satisfied", actor=actor,
                             reason="Reviewed the synthetic source and verified the email submission instructions.", expected_version=requirement.version)
        review = set_review_policy(db, opportunity_id=opp.id, review_policy="dual", actor=actor)
        for reviewer in reviewers:
            assign_reviewer(db, opportunity_id=opp.id, user_id=reviewer.id, actor_user_id=actor.id)
            add_comment(db, opportunity_id=opp.id, user_id=reviewer.id, body="Reviewed source, pricing and destination; synthetic bid can proceed.", validate_with_ai=False)
            complete_assignment(db, opportunity_id=opp.id, user_id=reviewer.id, action="approve_continue", recommendation="bid")
        db.refresh(review)
        finalize_approval(db, opportunity_id=opp.id, actor=actor, action="approve_to_bid", expected_version=review.version)
        generate_proposal(db, opportunity_id=opp.id, actor=actor, skip_ai=True, settings=configured)
        proposal = db.scalar(select(Proposal).where(Proposal.opportunity_id == opp.id))
        version = create_proposal_version(db, proposal_id=proposal.id, created_by=actor.email, actor_id=actor.id, status_after="ai_generated",
            sections=[{"section_key":"submission", "heading":"Submission instructions", "content":text, "requirement_ids":[requirement.id]}])
        check_proposal_coverage(db, opp.id, version.id, use_ai=False, settings=configured)
        proposal_path = artifact_dir / "proposal.docx"
        proposal_path.write_bytes(export_proposal_docx(db, proposal_version_id=version.id))
        pricing_path = artifact_dir / "pricing.xlsx"
        book = Workbook()
        book.active.append(["CLIN", "Quantity", "Unit price"])
        book.active.append(["0001", 2, 100])
        book.save(pricing_path)
        generate_submission_package(db, opportunity_id=opp.id, actor=actor, settings=configured)
        package = SubmissionPackage(files=[PackageFile("proposal.docx", role="proposal", local_path=str(proposal_path)),
                                          PackageFile("pricing.xlsx", role="pricing", local_path=str(pricing_path))],
            proposal_version_id=version.id, submission_method="email", recipient_email="release@example.test",
            pricing_rows=[{"clin":"0001", "quantity":2, "unit_price":100}], amendments_acknowledged=[])
        assemble_package(db, opportunity_id=opp.id, package=package, actor=actor)
        submission = db.scalar(select(Submission).where(Submission.opportunity_id == opp.id))
        ids = (opp.id, actor.id, proposal.id, submission.id)
        db.commit()
    return ids, package, configured


def run_release_lifecycle(engine, artifact_dir):
    from govcon.compliance.submission_preflight import run_submission_preflight
    from govcon.models import Proposal, Submission, User
    from govcon.proposals.service import finalize_proposal, record_submission_confirmation
    from govcon.submissions.checklist import generate_final_checklist
    ids, package, configured = prepared_lifecycle(engine, artifact_dir)
    opp_id, actor_id, proposal_id, submission_id = ids
    with Session(engine) as db:
        result = run_submission_preflight(db, opp_id, package, use_ai=False, settings=configured)
        assert result["ready"], result["items"]
        proposal = db.get(Proposal, proposal_id)
        actor = db.get(User, actor_id)
        finalize_proposal(db, opportunity_id=opp_id, action="APPROVE_FOR_SUBMISSION", actor=actor, expected_version=proposal.version)
        checklist = generate_final_checklist(db, opportunity_id=opp_id)
        assert checklist["overall"] == "ready", checklist
        submission = db.get(Submission, submission_id)
        record_submission_confirmation(db, opportunity_id=opp_id, actor=actor, expected_version=submission.version,
                                       confirmation_number="SYNTHETIC-RECEIPT-001", confirmation_notes="Synthetic human confirmation; no external submission was made.")
        assert submission.status in {"submitted", "confirmed"}
        db.commit()
    return ids


def test_full_release_lifecycle(upgraded_engine, tmp_path):
    run_release_lifecycle(upgraded_engine, tmp_path)


def test_release_package_failure_and_recovery(upgraded_engine, tmp_path):
    from govcon.compliance.submission_preflight import ReadinessBlocked, run_submission_preflight
    from govcon.models import Proposal, User
    from govcon.proposals.service import finalize_proposal
    ids, package, configured = prepared_lifecycle(upgraded_engine, tmp_path)
    opp_id, actor_id, proposal_id, _ = ids
    original = (tmp_path / "proposal.docx").read_bytes()
    (tmp_path / "proposal.docx").write_bytes(b"corrupted artifact")
    with Session(upgraded_engine) as db:
        result = run_submission_preflight(db, opp_id, package, use_ai=False, settings=configured)
        assert not result["ready"]
        with pytest.raises(ReadinessBlocked):
            finalize_proposal(db, opportunity_id=opp_id, action="APPROVE_FOR_SUBMISSION", actor=db.get(User, actor_id), expected_version=db.get(Proposal, proposal_id).version)
        db.commit()
    (tmp_path / "proposal.docx").write_bytes(original)
    with Session(upgraded_engine) as db:
        result = run_submission_preflight(db, opp_id, package, use_ai=False, settings=configured)
        assert result["ready"], result["items"]
        finalize_proposal(db, opportunity_id=opp_id, action="APPROVE_FOR_SUBMISSION", actor=db.get(User, actor_id), expected_version=db.get(Proposal, proposal_id).version)
        db.commit()


def test_concurrent_release_approvals_accept_one_version(upgraded_engine, tmp_path):
    from govcon.compliance.submission_preflight import run_submission_preflight
    from govcon.models import Proposal, User
    from govcon.proposals.service import finalize_proposal
    from govcon.proposals.versions import ProposalWorkflowError
    ids, package, configured = prepared_lifecycle(upgraded_engine, tmp_path)
    opp_id, actor_id, proposal_id, _ = ids
    with Session(upgraded_engine) as db:
        assert run_submission_preflight(db, opp_id, package, use_ai=False, settings=configured)["ready"]
        version = db.get(Proposal, proposal_id).version
        db.commit()
    barrier = Barrier(2)
    def approve():
        with Session(upgraded_engine) as db:
            actor = db.get(User, actor_id)
            barrier.wait(timeout=10)
            try:
                finalize_proposal(db, opportunity_id=opp_id, action="APPROVE_FOR_SUBMISSION", actor=actor, expected_version=version)
                db.commit()
                return True
            except ProposalWorkflowError:
                db.rollback()
                return False
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(approve) for _ in range(2)]
        assert sum(f.result(timeout=30) for f in futures) == 1
