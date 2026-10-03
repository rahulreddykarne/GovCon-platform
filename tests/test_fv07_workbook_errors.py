"""Quote workbook parsing preserves outputs and rejects corrupt uploads."""
import io
from decimal import Decimal
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from sqlalchemy import select
from test_sourcing_company import opportunity, user
from typer.testing import CliRunner

from govcon.cli import app
from govcon.models import SupplierQuote
from govcon.sourcing.records import parse_quote_table


def workbook_bytes(empty=False):
    workbook = Workbook()
    if not empty:
        workbook.active.append(["description", "part_number", "quantity", "unit", "unit_price"])
        workbook.active.append(["Gloves", "NG-100", 500, "PR", 2])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.parametrize("empty", [False, True])
def test_current_workbook_output(empty):
    lines = parse_quote_table(workbook_bytes(empty), "quote.xlsx")
    if empty:
        assert lines == []
    else:
        assert len(lines) == 1
        assert (lines[0]["part_number"], lines[0]["quantity"], lines[0]["unit_price"]) == ("NG-100", Decimal(500), Decimal(2))


@pytest.mark.parametrize("entry", ["web", "cli"])
def test_current_workbook_upload_caller_contracts(db, client, tmp_path, entry):
    opp = opportunity(db)
    actor, token = user(db)
    data = workbook_bytes()
    supplier = f"xlsx-{uuid4().hex}"
    if entry == "web":
        client.cookies.set("govcon_session", token)
        response = client.post(f"/workspace/{opp.id}/quotes", data={"supplier_name": supplier},
                               files={"quote_file": ("quote.xlsx", data, "application/octet-stream")})
        assert response.status_code == 303 and "notice=" in response.headers["location"]
    else:
        path = tmp_path / "quote.xlsx"
        path.write_bytes(data)
        result = CliRunner().invoke(app, ["sourcing", "add-quote", "--opportunity-id", str(opp.id),
                                        "--supplier", supplier, "--file", str(path), "--actor-email", actor.email])
        assert result.exit_code == 0 and "quote_id:" in result.output
    quote = db.scalar(select(SupplierQuote).where(SupplierQuote.opportunity_id == opp.id))
    assert quote.total_price == Decimal(1000) and quote.extraction_method == "xlsx"


def corrupt_workbook(kind):
    if kind == "not_zip":
        return b"not a ZIP workbook"
    buffer = io.BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        if kind == "bad_xml":
            with ZipFile(io.BytesIO(workbook_bytes())) as original:
                for name in original.namelist():
                    archive.writestr(name, b"<broken" if name == "[Content_Types].xml" else original.read(name))
    return buffer.getvalue()


@pytest.mark.parametrize("kind", ["not_zip", "empty_zip", "bad_xml"])
def test_invalid_workbooks_use_the_sourcing_error_contract(kind):
    from govcon.sourcing.records import SourcingError
    with pytest.raises(SourcingError, match="workbook"):
        parse_quote_table(corrupt_workbook(kind), "quote.xlsx")


@pytest.mark.parametrize("entry", ["web", "cli"])
def test_corrupt_upload_is_rejected_cleanly(db, client, tmp_path, entry):
    opp = opportunity(db)
    actor, token = user(db)
    supplier = f"corrupt-{uuid4().hex}"
    data = corrupt_workbook("not_zip")
    if entry == "web":
        client.cookies.set("govcon_session", token)
        response = client.post(f"/workspace/{opp.id}/quotes", data={"supplier_name": supplier},
                               files={"quote_file": ("quote.xlsx", data, "application/octet-stream")})
        assert response.status_code == 303 and "error=" in response.headers["location"]
    else:
        path = tmp_path / "quote.xlsx"
        path.write_bytes(data)
        result = CliRunner().invoke(app, ["sourcing", "add-quote", "--opportunity-id", str(opp.id),
                                        "--supplier", supplier, "--file", str(path), "--actor-email", actor.email])
        assert result.exit_code == 2 and "workbook" in result.output
    assert db.scalar(select(SupplierQuote.id).where(SupplierQuote.opportunity_id == opp.id)) is None


def test_loaded_read_only_workbook_is_closed(monkeypatch):
    from types import SimpleNamespace
    closed = []
    workbook = SimpleNamespace(active=SimpleNamespace(iter_rows=lambda **kw: iter([])), close=lambda: closed.append(True))
    monkeypatch.setattr("openpyxl.load_workbook", lambda *args, **kw: workbook)
    assert parse_quote_table(b"fake", "quote.xlsx") == []
    assert closed == [True]
