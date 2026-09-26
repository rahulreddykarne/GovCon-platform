"""Phase 7 tests: attachment ingestion, text extraction, AI analysis, and prompt registry.

Acceptance criteria from PHASE_07_ATTACHMENTS_AI_ANALYSIS.md:
1. Fixture PDF produces valid structured JSON.
2. Key values map to source references.
3. No API key = graceful warning.
4. Output is saved to ai_analyses, not directly over source fields.
"""

from __future__ import annotations

import hashlib
import json
import os
import uuid
import warnings
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from govcon.models import AIAnalysis, Opportunity, StoredFile

FIXTURE_PDF = Path(__file__).parent / "fixtures" / "solicitation_fixture.pdf"


def _uid() -> str:
    return uuid.uuid4().hex[:8]


# ── Helpers ──


def _make_opportunity(session: Session, **overrides) -> Opportunity:
    defaults = {
        "source": "sam",
        "source_id": f"test-phase7-{_uid()}",
        "title": "Test Solicitation for Phase 7",
        "status": "open",
        "raw": {"noticeId": "test-phase7-001"},
        "links": {},
        "psc_code": "6515",
        "naics_code": "339112",
        "set_aside_code": "SBA",
        "agency_path": "DEPT OF DEFENSE.DLA.DLA TROOP SUPPORT",
    }
    defaults.update(overrides)
    opp = Opportunity(**defaults)
    session.add(opp)
    session.flush()
    return opp


def _make_stored_file(session: Session, opp: Opportunity) -> StoredFile:
    """Create a StoredFile from the fixture PDF."""
    data = FIXTURE_PDF.read_bytes()
    from govcon.enrich.extract import extract_text

    result = extract_text(data, "application/pdf", "solicitation_fixture.pdf")
    sf = StoredFile(
        opportunity_id=opp.id,
        filename="solicitation_fixture.pdf",
        url=None,
        local_path=str(FIXTURE_PDF),
        mime_type="application/pdf",
        sha256=hashlib.sha256(data).hexdigest(),
        extracted_text=result.text,
        extraction_status=result.status,
        extraction_error=result.error,
    )
    session.add(sf)
    session.flush()
    return sf


# ── Text extraction tests ──


class TestTextExtraction:
    def test_pdf_extraction(self):
        from govcon.enrich.extract import extract_text

        data = FIXTURE_PDF.read_bytes()
        result = extract_text(data, "application/pdf", "test.pdf")
        assert result.status == "success"
        assert result.page_count and result.page_count >= 1
        assert "SPE4A6-26-Q-0042" in result.text
        assert "6515-01-519-8818" in result.text
        assert "50,000" in result.text

    def test_plain_text_extraction(self):
        from govcon.enrich.extract import extract_text

        data = b"This is a plain text solicitation document.\nNSN: 6515-01-519-8818"
        result = extract_text(data, "text/plain", "test.txt")
        assert result.status == "success"
        assert "6515-01-519-8818" in result.text

    def test_unsupported_mime_type(self):
        from govcon.enrich.extract import extract_text

        result = extract_text(b"binary", "application/zip", "test.zip")
        assert result.status == "unsupported"
        assert result.text is None

    def test_docx_extraction(self):
        from govcon.enrich.extract import extract_text

        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        import io

        doc = Document()
        doc.add_paragraph("Test solicitation DOCX content.")
        doc.add_paragraph("NSN: 6515-01-632-0167, Quantity: 100")
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()

        result = extract_text(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "test.docx")
        assert result.status == "success"
        assert "6515-01-632-0167" in result.text

    def test_xlsx_extraction(self):
        from govcon.enrich.extract import extract_text

        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl not installed")

        import io

        wb = Workbook()
        ws = wb.active
        ws.append(["NSN", "Quantity", "Unit"])
        ws.append(["6515-01-519-8818", 50000, "PR"])
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()

        result = extract_text(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "test.xlsx")
        assert result.status == "success"
        assert "6515-01-519-8818" in result.text

    def test_guess_mime_type(self):
        from govcon.enrich.extract import guess_mime_type

        assert guess_mime_type("test.pdf") == "application/pdf"
        assert guess_mime_type("test.docx") == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert guess_mime_type("test.xlsx") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert guess_mime_type("test.txt") == "text/plain"


# ── Attachment ingestion tests ──


class TestAttachmentIngestion:
    def test_process_local_file(self, upgraded_engine):
        from govcon.enrich.attachments import process_local_file

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"attach-local-{_uid()}")
            sf = process_local_file(session, opp, FIXTURE_PDF)
            session.commit()

            assert sf.id is not None
            assert sf.sha256 is not None
            assert sf.extraction_status == "success"
            assert "SPE4A6-26-Q-0042" in sf.extracted_text
            assert sf.filename == "solicitation_fixture.pdf"
            assert sf.mime_type == "application/pdf"

    def test_dedup_identical_content(self, upgraded_engine):
        from govcon.enrich.attachments import process_local_file

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"attach-dedup-{_uid()}")
            sf1 = process_local_file(session, opp, FIXTURE_PDF)
            sf2 = process_local_file(session, opp, FIXTURE_PDF)
            session.commit()
            assert sf1.id == sf2.id

    def test_sha256_computed(self, upgraded_engine):
        from govcon.enrich.attachments import process_local_file

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"attach-sha-{_uid()}")
            sf = process_local_file(session, opp, FIXTURE_PDF)
            session.commit()

            expected_sha = hashlib.sha256(FIXTURE_PDF.read_bytes()).hexdigest()
            assert sf.sha256 == expected_sha

    def test_extraction_status_tracked(self, upgraded_engine):
        from govcon.enrich.attachments import process_local_file

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"attach-status-{_uid()}")
            sf = process_local_file(session, opp, FIXTURE_PDF)
            session.commit()
            assert sf.extraction_status in ("success", "partial")


# ── AI output schema validation ──


class TestAISchemas:
    def test_valid_solicitation_analysis(self):
        from govcon.ai.schemas import validate_analysis_output

        data = {
            "summary": "DLA solicitation for surgical and examination gloves.",
            "items": [
                {
                    "description": "Surgical Gloves, Nitrile",
                    "nsn": "6515-01-519-8818",
                    "quantity": 50000,
                    "unit": "PR",
                    "source_refs": [{"page": 1, "section": "SECTION 2"}],
                }
            ],
            "delivery": {
                "location": "DLA Distribution Center, New Cumberland, PA",
                "delivery_days": 30,
                "source_refs": [{"page": 1, "section": "SECTION 3"}],
            },
            "eligibility": {
                "set_aside": "Total Small Business Set-Aside",
                "source_refs": [{"page": 1, "section": "SECTION 1"}],
            },
            "submission": {
                "method": "Email",
                "recipient_email": "DLATroopSupport@dla.mil",
                "deadline": "January 15, 2027, 2:00 PM Eastern Time",
                "timezone": "Eastern",
                "required_files": ["SF 1449", "Price Schedule", "Technical Capability Statement"],
                "source_refs": [{"page": 1, "section": "SECTION 4"}],
            },
            "key_dates": [
                {"label": "Response Deadline", "date": "2027-01-15T14:00:00", "timezone": "ET"}
            ],
            "risk_flags": [
                {
                    "category": "timeline",
                    "description": "Response deadline less than 30 days from issuance",
                    "source_refs": [{"page": 2, "section": "SECTION 10"}],
                }
            ],
            "clauses": ["FAR 52.212-1", "FAR 52.225-1"],
            "certifications": ["SAM.gov registration"],
            "country_of_origin_references": ["DFARS 252.225-7001"],
            "past_performance_requirements": [
                "Two references for similar contracts within last 3 years"
            ],
            "missing_information": [],
            "source_refs": [{"source_file_id": 1, "page": 1}],
        }
        result = validate_analysis_output("solicitation_analysis.v1", data)
        assert result.summary is not None
        assert len(result.items) >= 1
        assert result.items[0].nsn == "6515-01-519-8818"

    def test_minimal_valid_output(self):
        from govcon.ai.schemas import validate_analysis_output

        data = {"summary": "Minimal analysis."}
        result = validate_analysis_output("solicitation_analysis.v1", data)
        assert result.summary == "Minimal analysis."
        assert result.items == []

    def test_unknown_schema_raises(self):
        from govcon.ai.schemas import validate_analysis_output

        with pytest.raises(ValueError, match="unknown analysis schema"):
            validate_analysis_output("unknown.v99", {})

    def test_source_refs_in_items(self):
        from govcon.ai.schemas import validate_analysis_output

        data = {
            "items": [
                {
                    "description": "Widget",
                    "quantity": 100,
                    "source_refs": [
                        {"source_file_id": 1, "page": 3, "section": "CLIN 0001", "quote": "Quantity: 100"}
                    ],
                }
            ]
        }
        result = validate_analysis_output("solicitation_analysis.v1", data)
        assert result.items[0].source_refs[0].page == 3


# ── Prompt system tests ──


class TestPromptSystem:
    def test_shared_prompts_have_production_bodies(self):
        from govcon.prompting.loader import load_markdown_prompt

        prompt_root = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
        shared = prompt_root / "shared"

        for name in [
            "source_security_rules_v1",
            "no_fabrication_rules_v1",
            "evidence_rules_v1",
            "company_facts_policy_v1",
        ]:
            asset = load_markdown_prompt(shared / f"{name}.md")
            assert asset.metadata.get("status") == "active", f"{name} should be active"
            assert len(asset.body.strip()) > 50, f"{name} body is too short"
            assert "PLACEHOLDER" not in asset.body, f"{name} still has placeholder text"

    def test_solicitation_analysis_prompt_active(self):
        from govcon.prompting.loader import load_markdown_prompt

        prompt_root = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
        asset = load_markdown_prompt(prompt_root / "deepseek" / "solicitation_analysis_v1.md")
        assert asset.metadata.get("status") == "active"
        assert "solicitation analysis engine" in asset.body.lower()
        assert "PLACEHOLDER" not in asset.body

    def test_solicitation_prompt_includes_shared_fragments(self):
        from govcon.prompting.loader import load_markdown_prompt
        from govcon.prompting.renderer import render_system_prompt

        prompt_root = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
        asset = load_markdown_prompt(prompt_root / "deepseek" / "solicitation_analysis_v1.md")
        rendered = render_system_prompt(asset, prompt_root)
        assert "SOURCE SECURITY RULES" in rendered
        assert "NO-FABRICATION RULES" in rendered
        assert "EVIDENCE RULES" in rendered
        assert "solicitation analysis engine" in rendered.lower()

    def test_prompt_registry_sync(self, upgraded_engine):
        from govcon.prompting.registry import sync_prompts

        prompt_root = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
        with Session(upgraded_engine) as session:
            synced = sync_prompts(session, prompt_root)
            session.commit()
        assert "solicitation_analysis" in synced
        assert "source_security_rules" in synced

    def test_prompt_load_from_disk(self):
        from govcon.prompting.registry import load_prompt_from_disk

        prompt_root = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
        asset = load_prompt_from_disk(prompt_root, "solicitation_analysis")
        assert asset.name == "solicitation_analysis"
        assert asset.version == "v1"
        assert len(asset.body) > 100

    def test_no_hardcoded_prompts_in_business_services(self):
        """§43.1: no production AI task prompt is hardcoded inside a business-service module."""
        import ast

        service_dirs = [
            Path(__file__).parent.parent / "src" / "govcon" / "enrich",
            Path(__file__).parent.parent / "src" / "govcon" / "intelligence",
        ]
        for sdir in service_dirs:
            for py_file in sdir.rglob("*.py"):
                source = py_file.read_text()
                tree = ast.parse(source, str(py_file))
                for node in ast.walk(tree):
                    if isinstance(node, ast.Constant) and isinstance(node.value, str):
                        val = node.value
                        if len(val) > 200 and "you are" in val.lower():
                            pytest.fail(
                                f"potential hardcoded prompt in {py_file.name}: "
                                f"{val[:100]}..."
                            )


# ── AI provider tests ──


class TestAIProvider:
    def test_no_api_key_graceful_warning(self, upgraded_engine):
        """Acceptance: No API key = graceful warning."""
        from govcon.enrich.summarize import AnalysisWarning, run_solicitation_analysis

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"nokey-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with warnings.catch_warnings(record=True) as w:
                warnings.simplefilter("always")
                os.environ.pop("DEEPSEEK_API_KEY", None)
                os.environ.pop("ANTHROPIC_API_KEY", None)
                os.environ.pop("OPENAI_API_KEY", None)
                from govcon.config import get_settings
                get_settings.cache_clear()

                result = run_solicitation_analysis(session, opp)
                assert result is None
                warning_messages = [str(x.message) for x in w if issubclass(x.category, AnalysisWarning)]
                assert any("No AI provider configured" in m for m in warning_messages)

    def test_provider_factory_no_key(self):
        from govcon.ai.providers import NoProviderConfigured, get_provider
        from govcon.config import Settings

        settings = Settings(deepseek_api_key=None, database_url="postgresql+psycopg://x@y/z")
        with pytest.raises(NoProviderConfigured):
            get_provider(settings)

    def test_deepseek_provider_init(self):
        from govcon.ai.providers.deepseek import DeepSeekProvider

        p = DeepSeekProvider(api_key="test-key", model="deepseek-flash")
        assert p.name == "deepseek"
        assert p._model == "deepseek-flash"


# ── Solicitation analysis integration ──


class TestSolicitationAnalysis:
    def _mock_deepseek_response(self):
        """Return a valid solicitation analysis JSON as if DeepSeek produced it."""
        return {
            "summary": "DLA solicitation for nitrile and latex-free gloves.",
            "items": [
                {
                    "description": "Surgical Gloves, Nitrile, Powder-Free",
                    "nsn": "6515-01-519-8818",
                    "quantity": 50000,
                    "unit": "PR",
                    "manufacturer": "Ansell or Equal",
                    "part_number": "93-850",
                    "source_refs": [
                        {"source_file_id": 1, "page": 1, "section": "SECTION 2", "quote": "CLIN 0001 - Surgical Gloves"}
                    ],
                },
                {
                    "description": "Examination Gloves, Latex-Free",
                    "nsn": "6515-01-632-0167",
                    "quantity": 100000,
                    "unit": "EA",
                    "source_refs": [
                        {"source_file_id": 1, "page": 1, "section": "SECTION 2", "quote": "CLIN 0002 - Examination Gloves"}
                    ],
                },
            ],
            "delivery": {
                "location": "DLA Distribution Center, New Cumberland, PA 17070",
                "delivery_days": 30,
                "source_refs": [{"source_file_id": 1, "page": 1, "section": "SECTION 3"}],
            },
            "eligibility": {
                "set_aside": "Total Small Business Set-Aside (SBA)",
                "source_refs": [{"source_file_id": 1, "page": 1, "section": "SECTION 1"}],
            },
            "submission": {
                "method": "Email",
                "recipient_email": "DLATroopSupport@dla.mil",
                "deadline": "January 15, 2027, 2:00 PM Eastern Time",
                "timezone": "Eastern",
                "required_files": ["SF 1449", "Price Schedule", "Technical Capability Statement"],
                "source_refs": [{"source_file_id": 1, "page": 1, "section": "SECTION 4"}],
            },
            "key_dates": [
                {"label": "Response Deadline", "date": "2027-01-15T14:00:00", "timezone": "ET"}
            ],
            "evaluation_factors": [
                {"name": "Technical Capability", "weight": "pass/fail"},
                {"name": "Price", "weight": "lowest price"},
                {"name": "Past Performance", "weight": "acceptable/unacceptable"},
            ],
            "certifications": [
                "SAM.gov registration (active)",
                "FAR 52.212-3 Representations and Certifications",
                "Buy American Act compliance (FAR 52.225-1)",
            ],
            "country_of_origin_references": ["DFARS 252.225-7001"],
            "past_performance_requirements": [
                "At least two past performance references for similar contracts within last 3 years"
            ],
            "clauses": [
                "FAR 52.212-1",
                "FAR 52.212-4",
                "FAR 52.219-6",
                "FAR 52.225-1",
                "DFARS 252.225-7001",
            ],
            "risk_flags": [
                {
                    "category": "timeline",
                    "description": "Response deadline less than 30 days from issuance",
                    "severity": "medium",
                    "source_refs": [{"source_file_id": 1, "page": 2, "section": "SECTION 10"}],
                },
                {
                    "category": "delivery",
                    "description": "Aggressive delivery timeline for the quantity requested",
                    "severity": "medium",
                    "source_refs": [{"source_file_id": 1, "page": 2, "section": "SECTION 10"}],
                },
            ],
            "pricing_structure": {
                "format": "Unit price per item, extended total per CLIN",
                "source_refs": [{"source_file_id": 1, "page": 1, "section": "SECTION 4"}],
            },
            "conflicts": [],
            "missing_information": [],
            "source_refs": [{"source_file_id": 1, "page": 1}],
        }

    def test_fixture_pdf_produces_valid_json(self, upgraded_engine):
        """Acceptance: Fixture PDF produces valid structured JSON."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_response = self._mock_deepseek_response()
        mock_result = DeepSeekResult(
            content=json.dumps(mock_response),
            model="deepseek-flash",
            usage={"prompt_tokens": 1500, "completion_tokens": 800, "total_tokens": 2300},
            latency_ms=2500,
            finish_reason="stop",
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-fixture-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                analysis = run_solicitation_analysis(session, opp, force=True)
                session.commit()

            assert analysis is not None
            assert analysis.output_json is not None
            output = analysis.output_json
            assert output["summary"] is not None
            assert len(output["items"]) == 2
            assert output["items"][0]["nsn"] == "6515-01-519-8818"
            assert output["items"][0]["quantity"] == 50000

    def test_key_values_map_to_source_refs(self, upgraded_engine):
        """Acceptance: Key values map to source references."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_response = self._mock_deepseek_response()
        mock_result = DeepSeekResult(
            content=json.dumps(mock_response),
            model="deepseek-flash",
            usage={},
            latency_ms=1000,
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-refs-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                analysis = run_solicitation_analysis(session, opp, force=True)
                session.commit()

            output = analysis.output_json
            item_refs = output["items"][0].get("source_refs", [])
            assert len(item_refs) >= 1
            assert item_refs[0]["page"] is not None
            assert item_refs[0]["section"] is not None

            delivery_refs = output["delivery"].get("source_refs", [])
            assert len(delivery_refs) >= 1

            submission_refs = output["submission"].get("source_refs", [])
            assert len(submission_refs) >= 1

    def test_output_saved_to_ai_analyses(self, upgraded_engine):
        """Acceptance: Output is saved to ai_analyses, not directly over source fields."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_response = self._mock_deepseek_response()
        mock_result = DeepSeekResult(
            content=json.dumps(mock_response),
            model="deepseek-flash",
            usage={},
            latency_ms=1000,
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-persist-{_uid()}")
            _make_stored_file(session, opp)
            original_title = opp.title
            original_description = opp.description
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                analysis = run_solicitation_analysis(session, opp, force=True)
                session.commit()

            assert analysis is not None
            assert analysis.analysis_type == "solicitation_summary"
            assert analysis.schema_version == "solicitation_analysis.v1"

            session.refresh(opp)
            assert opp.title == original_title
            assert opp.description == original_description

            stored = session.execute(
                select(AIAnalysis).where(AIAnalysis.id == analysis.id)
            ).scalar_one()
            assert stored.output_json["summary"] is not None
            assert stored.prompt_name == "solicitation_analysis"
            assert stored.prompt_hash is not None
            assert stored.context_manifest is not None

    def test_analysis_records_prompt_metadata(self, upgraded_engine):
        """§43.2: Every AI run must be reproducible from recorded metadata."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_response = self._mock_deepseek_response()
        mock_result = DeepSeekResult(
            content=json.dumps(mock_response),
            model="deepseek-flash",
            usage={"prompt_tokens": 100, "completion_tokens": 50, "total_tokens": 150},
            latency_ms=500,
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-meta-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                analysis = run_solicitation_analysis(session, opp, force=True)
                session.commit()

            assert analysis.provider == "deepseek"
            assert analysis.model == "deepseek-flash"
            assert analysis.prompt_name == "solicitation_analysis"
            assert analysis.prompt_version == "v1"
            assert analysis.prompt_hash is not None
            assert len(analysis.prompt_hash) == 64
            assert analysis.schema_version == "solicitation_analysis.v1"
            assert analysis.generation_settings is not None
            assert analysis.input_snapshot_hash is not None
            assert analysis.context_manifest is not None
            assert "files" in analysis.context_manifest

    def test_existing_analysis_not_overwritten(self, upgraded_engine):
        """Re-running without force does not create a second analysis."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_response = self._mock_deepseek_response()
        mock_result = DeepSeekResult(
            content=json.dumps(mock_response),
            model="deepseek-flash",
            usage={},
            latency_ms=1000,
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-nodup-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                a1 = run_solicitation_analysis(session, opp, force=True)
                session.commit()
                a2 = run_solicitation_analysis(session, opp, force=False)
                assert a2.id == a1.id

    def test_malformed_json_response_returns_none(self, upgraded_engine):
        """Malformed structured output fails closed."""
        from govcon.ai.providers.deepseek import DeepSeekResult
        from govcon.enrich.summarize import run_solicitation_analysis

        mock_result = DeepSeekResult(
            content="this is not valid json {{{",
            model="deepseek-flash",
            usage={},
            latency_ms=100,
        )

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-bad-json-{_uid()}")
            _make_stored_file(session, opp)
            session.flush()

            with patch("govcon.enrich.summarize.get_provider") as mock_get:
                mock_provider = MagicMock()
                mock_provider.complete.return_value = mock_result
                mock_provider.name = "deepseek"
                mock_get.return_value = mock_provider

                result = run_solicitation_analysis(session, opp, force=True)
                assert result is None

    def test_no_files_returns_none(self, upgraded_engine):
        """No extracted text = no analysis."""
        from govcon.enrich.summarize import run_solicitation_analysis

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"analysis-nofiles-{_uid()}")
            session.flush()
            result = run_solicitation_analysis(session, opp)
            assert result is None


# ── DeepSeek provider JSON parsing ──


class TestDeepSeekParsing:
    def test_parse_json_response(self):
        from govcon.ai.providers.deepseek import DeepSeekResult, parse_json_response

        result = DeepSeekResult(
            content='{"summary": "test"}',
            model="deepseek-flash",
        )
        parsed = parse_json_response(result)
        assert parsed == {"summary": "test"}

    def test_parse_markdown_wrapped_json(self):
        from govcon.ai.providers.deepseek import DeepSeekResult, parse_json_response

        result = DeepSeekResult(
            content='```json\n{"summary": "test"}\n```',
            model="deepseek-flash",
        )
        parsed = parse_json_response(result)
        assert parsed == {"summary": "test"}

    def test_parse_invalid_json_raises(self):
        from govcon.ai.providers.deepseek import DeepSeekResult, parse_json_response

        result = DeepSeekResult(
            content="not json at all",
            model="deepseek-flash",
        )
        with pytest.raises(json.JSONDecodeError):
            parse_json_response(result)


# ── Context manifest ──


class TestContextManifest:
    def test_context_manifest_structure(self, upgraded_engine):
        from govcon.enrich.summarize import _build_context_manifest

        with Session(upgraded_engine) as session:
            opp = _make_opportunity(session, source_id=f"manifest-{_uid()}")
            sf = _make_stored_file(session, opp)
            session.flush()
            manifest = _build_context_manifest(opp, [sf])
            assert manifest["opportunity_id"] == opp.id
            assert len(manifest["files"]) == 1
            assert manifest["files"][0]["sha256"] == sf.sha256
            assert manifest["files"][0]["file_id"] == sf.id
