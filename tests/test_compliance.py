"""Phase 9 tests: high-reliability compliance subsystem (MASTER_SPEC §15.22).

DB-backed tests run against the Compose/CI PostgreSQL. External AI providers
are mocked; no live DeepSeek/Claude/JEV key is needed. JEV routing uses the
Phase 8 rule fallback when ``JEV_API_KEY`` is unset.
"""

from __future__ import annotations

import ast
import hashlib
import json
import shutil
from datetime import UTC, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session
from typer.testing import CliRunner

from govcon.ai.providers.deepseek import DeepSeekResult
from govcon.compliance import deterministic as det
from govcon.compliance.deterministic import PackageFile, SubmissionPackage
from govcon.compliance.matrix import (
    ComplianceInvariantError,
    StatusDecision,
    ValidationInputs,
    add_evidence,
    apply_decision,
    compliance_matrix,
    decide_status,
    override_requirement,
    record_human_verification,
)
from govcon.compliance.pipeline import run_compliance_pipeline
from govcon.models import (
    AIAnalysis,
    AuditEvent,
    ClauseLibraryEntry,
    ComplianceFinding,
    ComplianceRun,
    DecisionRun,
    Opportunity,
    Proposal,
    ProposalSection,
    ProposalVersion,
    Pursuit,
    Requirement,
    StoredFile,
    Submission,
)

FIXTURES = Path(__file__).parent / "fixtures" / "compliance"
DLA = FIXTURES / "dla_base_amendment"
QA = FIXTURES / "qa_conflict_injection"
TABLE = FIXTURES / "table_embedded_packaging"
PROMPT_ROOT = Path(__file__).parent.parent / "src" / "govcon" / "prompts"
COMPLIANCE_PROMPTS = (
    "requirement_extraction_a", "requirement_extraction_b", "requirement_reconciliation", "compliance_validator",
    "contradiction_detection", "compliance_red_team", "amendment_analysis", "proposal_coverage", "submission_preflight_ai",
)
DEFAULTS = {
    "requirement_extraction_a": {"requirements": []},
    "requirement_extraction_b": {"requirements": []},
    "requirement_reconciliation": {"groups": []},
    "compliance_validator": {"validations": []},
    "contradiction_detection": {"conflicts": []},
    "compliance_red_team": {"findings": []},
    "amendment_analysis": {"material": None, "changes": []},
    "proposal_coverage": {"coverage": []},
    "submission_preflight_ai": {"status": "NEEDS_REVIEW", "issues": [], "unresolved": []},
}
NOW = datetime(2026, 10, 1, 12, 0, tzinfo=UTC)


class FakeProvider:
    """Stands in for DeepSeek; responses are keyed by prompt name (``purpose``)."""

    name = "deepseek"

    def __init__(self, handlers: dict | None = None) -> None:
        self.handlers = handlers or {}
        self.calls: list[dict] = []

    def complete(self, *, system_prompt, user_prompt, purpose, model=None, **_):
        self.calls.append({"purpose": purpose, "system": system_prompt, "user": user_prompt})
        handler = self.handlers.get(purpose, DEFAULTS.get(purpose))
        payload = handler(user_prompt) if callable(handler) else handler
        content = payload if isinstance(payload, str) else json.dumps(payload)
        return DeepSeekResult(content=content, model=model or "deepseek-flash", usage={"total_tokens": 10}, latency_ms=3)


def _uid() -> str:
    return uuid4().hex[:10]


def _opp(session: Session, **overrides) -> Opportunity:
    data = {
        "source": "sam",
        "source_id": f"phase9-{_uid()}",
        "title": "Phase 9 compliance fixture",
        "status": "open",
        "psc_code": "6515",
        "naics_code": "339113",
        "set_aside_code": "SBA",
        "response_deadline": datetime(2026, 10, 15, 18, 0, tzinfo=UTC),
        "raw": {"fixture": True},
        "links": {},
    }
    data.update(overrides)
    row = Opportunity(**data)
    session.add(row)
    session.flush()
    return row


def _file(session: Session, opp: Opportunity, text: str | None, filename: str, *, mime: str = "text/plain", status: str = "success", sha: str | None = None, url: str | None = None, error: str | None = None) -> StoredFile:
    row = StoredFile(classification="PUBLIC", source_origin="synthetic_test_fixture",
        opportunity_id=opp.id,
        filename=filename,
        url=url or f"https://example.test/{_uid()}/{filename}",
        mime_type=mime,
        sha256=sha or hashlib.sha256(f"{filename}{text}{_uid()}".encode()).hexdigest(),
        extracted_text=text,
        extraction_status=status,
        extraction_error=error,
        downloaded_at=datetime.now(UTC),
    )
    session.add(row)
    session.flush()
    return row


def _remap(path: Path, ids: dict[int, int]) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    for item in data["requirements"]:
        if item.get("source_file_id") in ids:
            item["source_file_id"] = ids[item["source_file_id"]]
    return data


def _dla_setup(session: Session, *, with_amendment: bool = False, **opp):
    o = _opp(session, **opp)
    f1 = _file(session, o, (DLA / "solicitation.txt").read_text(encoding="utf-8"), "SPE2DM-26-Q-0412_solicitation.txt")
    f2 = _file(session, o, (DLA / "price_schedule.txt").read_text(encoding="utf-8"), "Attachment_2_Price_Schedule.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ids = {1: f1.id, 2: f2.id}
    if with_amendment:
        ids[3] = _add_amendment(session, o).id
    return o, ids


def _add_amendment(session: Session, opp: Opportunity) -> StoredFile:
    return _file(session, opp, (DLA / "amendment_0001.txt").read_text(encoding="utf-8"), "SPE2DM-26-Q-0412_amendment_0001.txt")


def _dla_provider(ids: dict[int, int], extra: dict | None = None) -> FakeProvider:
    def passes(label: str):
        def handler(user_prompt: str):
            data = _remap(DLA / f"base_pass_{label}.json", ids)
            if 3 in ids and f"file_id={ids[3]} " in user_prompt:
                data["requirements"] += _remap(DLA / f"amendment_pass_{label}.json", ids)["requirements"]
            return data
        return handler

    handlers = {"requirement_extraction_a": passes("a"), "requirement_extraction_b": passes("b")}
    handlers.update(extra or {})
    return FakeProvider(handlers)


def _run(session, opp, provider, **kwargs):
    kwargs.setdefault("company_facts", {"sam_registration_status": "active", "sam_expiration_date": "2027-06-30"})
    kwargs.setdefault("now", NOW)
    with patch("govcon.ai.structured.get_provider", return_value=provider):
        return run_compliance_pipeline(session, opp.id, **kwargs)


def _req(session, opp, text: str) -> Requirement:
    rows = session.scalars(select(Requirement).where(Requirement.opportunity_id == opp.id)).all()
    hits = [r for r in rows if text.lower() in f"{r.requirement_text} {r.source_quote or ''}".lower()]
    assert hits, f"no requirement containing {text!r}"
    hits.sort(key=lambda r: (0 if r.independently_confirmed else 1, r.id))
    return hits[0]


def _user(session, role: str):
    from govcon.collaboration.users import invite_user

    return invite_user(session, email=f"{role}-{_uid()}@example.test", display_name=role, password="correct horse battery", role=role)


# Mocked AI here sends PROPRIETARY data; the default policy blocks that.
pytestmark = pytest.mark.usefixtures("allow_proprietary_ai")


@pytest.fixture(autouse=True)
def independent_pass_b(monkeypatch):
    """Pass B on a different model, so A/B agreement counts as independent confirmation."""
    from govcon.config import get_settings

    monkeypatch.setenv("COMPLIANCE_PASS_B_MODEL", "deepseek-v4-pro")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


@pytest.fixture()
def session(upgraded_engine):
    from govcon.compliance.clauses import seed_clause_library

    with Session(upgraded_engine) as s:
        seed_clause_library(s)
        from govcon.config import get_settings
        from govcon.prompting.registry import sync_prompts
        sync_prompts(s, get_settings().resolved_prompt_root())
        yield s
        s.rollback()


# ── AC1: document inventory + source attribution ──


def test_inventory_records_every_field_and_detects_problems(session) -> None:
    from govcon.compliance.inventory import build_document_inventory

    o = _opp(session, links={"resourceLinks": [{"url": "https://sam.test/att9.pdf", "name": "Attachment 9"}]})
    base = _file(session, o, "SOLICITATION W1 RFQ\nSee Attachment 4 for drawings.", "RFQ_solicitation.txt", sha="a" * 64)
    _file(session, o, "same bytes", "copy_of_rfq.txt", sha="a" * 64)
    _file(session, o, "v1 of the SOW", "SOW.txt", sha="b" * 64)
    _file(session, o, "v2 of the SOW", "SOW.txt", sha="c" * 64)
    _file(session, o, None, "scan.pdf", mime="application/pdf", status="partial", error="PDF contained no extractable text (may need OCR)")
    _file(session, o, None, "broken.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", status="error", error="bad zip")
    _file(session, o, "Amendment 0003 to the solicitation", "amendment_0003.txt")

    inventory, run = build_document_inventory(session, o.id)
    codes = {w.code for w in inventory.warnings}
    assert {"duplicate_file", "same_filename_different_hash", "unreadable_pages", "extraction_failed", "amendment_sequence_gap", "newer_amendment", "missing_expected_attachment", "possibly_missing_referenced_attachment"} <= codes
    assert run.status == "incomplete" and run.run_type == "document_inventory"
    doc = next(d for d in run.output_json["documents"] if d["file_id"] == base.id)
    for key in ("filename", "url", "sha256", "downloaded_at", "snapshot_id", "document_type", "page_count", "text_extraction_status", "table_extraction_status", "ocr_needed"):
        assert key in doc
    assert doc["document_type"] == "solicitation"
    assert next(d for d in run.output_json["documents"] if d["filename"] == "scan.pdf")["ocr_needed"] is True
    findings = session.scalars(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id)).all()
    assert any(f.finding_type == "inventory_extraction_failed" and f.blocks_submission for f in findings)


def test_pipeline_requirements_are_source_backed(session) -> None:
    o, ids = _dla_setup(session)
    result = _run(session, o, _dla_provider(ids))
    reqs = session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id)).all()
    assert len(reqs) >= 12
    for r in reqs:
        if r.source_file_id is None or not r.source_quote:
            assert r.status == "needs_review", r.requirement_text
        else:
            assert r.source_refs, r.requirement_text
            assert r.extraction_pass
    sam = _req(session, o, "System for Award Management (SAM) at the time")
    assert sam.source_file_id == ids[1] and sam.source_section == "SECTION 3 EVALUATION"
    assert "Offerors must be registered" in sam.source_quote
    assert result["inventory"]["complete"] is True
    analyses = session.scalars(select(AIAnalysis).where(AIAnalysis.opportunity_id == o.id, AIAnalysis.prompt_name == "requirement_extraction_a")).all()
    assert analyses and analyses[0].prompt_hash and analyses[0].schema_version == "requirement_extraction.v1"
    assert analyses[0].context_manifest["strategy"] == "A" and analyses[0].context_manifest["files"]


# ── AC2 / AC3: independent passes + reconciliation; single-pass retained ──


def test_independent_passes_reconcile_and_single_pass_is_kept(session) -> None:
    o, ids = _dla_setup(session)
    provider = _dla_provider(ids)
    _run(session, o, provider)
    purposes = [c["purpose"] for c in provider.calls]
    assert "requirement_extraction_a" in purposes and "requirement_extraction_b" in purposes
    a_call = next(c for c in provider.calls if c["purpose"] == "requirement_extraction_a")
    b_call = next(c for c in provider.calls if c["purpose"] == "requirement_extraction_b")
    assert a_call["system"] != b_call["system"], "passes must use different prompts"
    assert a_call["user"] != b_call["user"], "passes must use different context strategies"
    run_types = {r.run_type for r in session.scalars(select(ComplianceRun).where(ComplianceRun.opportunity_id == o.id)).all()}
    assert {"extraction_pass_a", "extraction_pass_b", "extraction_pass_d", "requirement_reconciliation"} <= run_types

    deadline = _req(session, o, "October 15, 2026")
    assert deadline.independently_confirmed is True
    page_limit = _req(session, o, "shall not exceed 10 pages")
    assert page_limit.independently_confirmed is False
    assert "B" in page_limit.reconciliation["found_by"] and "A" not in page_limit.reconciliation["found_by"]
    assert "single_pass" in page_limit.reconciliation["flags"]
    references = _req(session, o, "three past performance references")
    assert "ab_disagreement" in references.reconciliation["flags"]
    assert references.status == "needs_review"


def test_passes_on_the_same_model_are_not_independent_confirmation(session, monkeypatch) -> None:
    from govcon.config import get_settings

    monkeypatch.delenv("COMPLIANCE_PASS_B_MODEL", raising=False)
    get_settings.cache_clear()
    o, ids = _dla_setup(session)
    result = _run(session, o, _dla_provider(ids))
    passes = result["extraction"]["passes"]
    assert (passes["A"]["provider"], passes["A"]["model"]) == (passes["B"]["provider"], passes["B"]["model"])
    deadline = _req(session, o, "October 15, 2026")
    assert set(deadline.reconciliation["found_by"]) >= {"A", "B"}
    assert deadline.independently_confirmed is False
    assert "ab_same_model" in deadline.reconciliation["flags"]


def test_reconciler_never_merges_conflicting_values_or_drops_candidates() -> None:
    from govcon.compliance.records import Candidate
    from govcon.compliance.reconciler import reconcile

    a = Candidate("A-1", "A", "Deliver within 30 days", key_values={"delivery_days": 30}, supporting_quote="Delivery shall be made within 30 days.", source_file_id=1, citation_verified=True)
    b = Candidate("B-1", "B", "Deliver within 20 days", key_values={"delivery_days": 20}, supporting_quote="Delivery shall be made within 20 days.", source_file_id=2, citation_verified=True)
    c = Candidate("B-2", "B", "Completely different requirement about marking labels", supporting_quote="Labels shall be affixed.", source_file_id=1, citation_verified=True)
    out = reconcile([a, b, c], merge_threshold=0.72, duplicate_threshold=0.4)
    assert len(out) == 3
    assert sorted(c.candidate_id for group in out for c in group.candidates) == ["A-1", "B-1", "B-2"]


# ── AC4: deterministic validators ──


def test_deterministic_validators_pass_fail_unknown() -> None:
    pkg = SubmissionPackage(
        files=[
            PackageFile("technical.pdf", size_bytes=1_000, role="proposal", page_count=12),
            PackageFile("SF1449.pdf", size_bytes=1_000, role="form", form_id="SF 1449", signed=True),
            PackageFile("SF30.pdf", size_bytes=1_000, role="acknowledgment", form_id="SF 30", signed=None),
            PackageFile("prices.xlsx", size_bytes=30 * 1024 * 1024, role="pricing"),
        ],
        recipient_email="wrong@dla.mil",
        amendments_acknowledged=[],
        pricing_rows=[{"clin": "0001", "quantity": 500, "unit_price": 4.2}, {"clin": "0002", "quantity": 150, "unit_price": None}],
    )
    cases = [
        (det.deadline_not_passed(NOW - timedelta(hours=1), NOW), "fail"),
        (det.deadline_not_passed(NOW + timedelta(days=3), NOW), "pass"),
        (det.deadline_not_passed(None, NOW), "unknown"),
        (det.deadline_timezone_consistent({"response_deadline_date": "October 15, 2026", "response_deadline_time": "2:00 PM", "deadline_timezone": "ET"}, datetime(2026, 10, 15, 18, 0, tzinfo=UTC)), "pass"),
        (det.deadline_timezone_consistent({"response_deadline_date": "October 15, 2026", "response_deadline_time": "2:00 PM", "deadline_timezone": "CT"}, datetime(2026, 10, 15, 18, 0, tzinfo=UTC)), "fail"),
        (det.deadline_timezone_consistent({"response_deadline_date": "October 15, 2026"}, None), "unknown"),
        (det.required_files_present(["SF 1449", "SF 30"], pkg), "pass"),
        (det.required_forms_present(["DD 1155"], pkg), "fail"),
        (det.required_files_present(["SF 1449"], None), "unknown"),
        (det.signatures_confirmed(pkg, ["SF 1449"]), "pass"),
        (det.signatures_confirmed(pkg, ["SF 30"]), "unknown"),
        (det.amendments_acknowledged(["0001"], pkg), "fail"),
        (det.amendments_acknowledged([], pkg), "pass"),
        (det.page_count_within_limit(10, pkg), "fail"),
        (det.page_count_within_limit(10, None), "unknown"),
        (det.pricing_rows_populated(pkg), "fail"),
        (det.clins_accounted({"0001": 500, "0003": 5}, pkg), "fail"),
        (det.quantities_accounted({"0001": 500, "0002": 200}, pkg), "fail"),
        (det.file_types_allowed(["PDF"], pkg), "fail"),
        (det.filenames_match(["technical.pdf"], pkg), "pass"),
        (det.file_size_within_limit(25, pkg), "fail"),
        (det.recipient_matches("jane.buyer@dla.mil", pkg), "fail"),
        (det.portal_matches("DIBBS", pkg), "unknown"),
        (det.sam_registration_known({}, None), "unknown"),
        (det.sam_registration_known({"sam_registration_status": "expired"}, None), "fail"),
        (det.set_aside_matches("SBA", {}), "unknown"),
        (det.set_aside_matches("SBA", {"socioeconomic": {"small_business": False}}), "fail"),
        (det.set_aside_matches("SBA", {"socioeconomic": {"small_business": True}}), "pass"),
        (det.delivery_date_arithmetic(30, {"lead_time_days": 12}), "unknown"),
        (det.delivery_date_arithmetic(20, {"lead_time_days": 12, "transit_days": 10}), "fail"),
        (det.delivery_date_arithmetic(30, {"lead_time_days": 12, "transit_days": 5}), "pass"),
        (det.margin_arithmetic(90.0, 100.0), "fail"),
        (det.margin_arithmetic(120.0, 100.0, 10.0), "pass"),
        (det.margin_arithmetic(None, 100.0), "unknown"),
    ]
    for result, expected in cases:
        assert result.status == expected, (result.validator, result.reason)
        assert result.validator_version == det.VALIDATOR_VERSION
        assert set(result.as_dict()) == {"validator", "status", "reason", "evidence", "validator_version"}


# ── AC5: deterministic failures cannot be overridden silently ──


def test_ai_cannot_override_deterministic_failure(session) -> None:
    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids))
    recipient = _req(session, o, "jane.buyer@dla.mil")
    evidence = add_evidence(session, requirement_id=recipient.id, evidence_type="submission_package", verification_method="submission_preflight", verification_status="verified", description="package manifest")
    validator = {"validations": [{"requirement_id": recipient.id, "status": "SATISFIED", "reason": "looks fine", "evidence_refs": [{"evidence_id": evidence.id}], "confidence": 0.99}]}
    package = SubmissionPackage(files=[PackageFile("quote.pdf", role="proposal")], recipient_email="someone.else@dla.mil")
    repeated = _run(session, o, _dla_provider(ids, {"compliance_validator": validator}), package=package, force=True)
    session.refresh(recipient)
    assert recipient.status in {"missing", "needs_review"}
    assert recipient.validation["deterministic_fail"] is True
    assert "blocked_ai_claims" in recipient.validation, repeated.get("warnings")
    assert any("deterministic validator(s) failed" in c for c in recipient.validation["blocked_ai_claims"])
    finding = session.scalar(select(ComplianceFinding).where(ComplianceFinding.requirement_id == recipient.id, ComplianceFinding.finding_type == "ai_claim_blocked"))
    assert finding is not None and finding.status == "open"

    owner = _user(session, "owner")
    with pytest.raises(ComplianceInvariantError):
        override_requirement(session, requirement_id=recipient.id, status="satisfied", actor=owner, reason="Contracting officer confirmed alternate address.", expected_version=recipient.version)


# ── AC6: unknown stays distinct ──


def test_unknown_is_distinct_from_missing_and_satisfied(session) -> None:
    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids), supplier={"lead_time_days": 12})
    delivery = _req(session, o, "within 30 calendar days")
    assert delivery.status == "unknown"
    assert "transit time is not" in delivery.status_reason
    set_aside = _req(session, o, "small business under NAICS 339113")
    assert set_aside.status == "unknown"
    base = ValidationInputs(requirement_type="certification", mandatory=True, severity="high", has_source_location=True, current_status="unreviewed", stale=False)
    assert decide_status(base).status == "unknown"
    base.deterministic = [{"validator": "x", "status": "unknown", "reason": "not known"}]
    assert decide_status(base).status == "unknown"


# ── AC7: clause library ──


def test_clause_references_map_to_library(session) -> None:
    from govcon.compliance.clauses import match_references
    from govcon.compliance.text import find_clause_references

    assert session.scalar(select(ClauseLibraryEntry).where(ClauseLibraryEntry.clause_number == "52.204-7")).title == "System for Award Management"
    o, ids = _dla_setup(session)
    result = _run(session, o, _dla_provider(ids))
    run = session.get(ComplianceRun, result["clauses"]["run_id"])
    known = {r["clause"]: r for r in run.output_json["references"]}
    assert known["FAR 52.204-7"]["known"] and known["DFARS 252.225-7001"]["known"]
    assert known["DLAD 52.211-9010"]["known"] is False
    linked = session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id, Requirement.clause_library_id.is_not(None))).all()
    assert linked and all(r.validation["clause"]["verification_questions"] for r in linked)
    unknown = session.scalar(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id, ComplianceFinding.finding_type == "unknown_clause"))
    assert unknown is not None and "52.211-9010" in unknown.description

    library = {(e.clause_family, e.clause_number): e for e in session.scalars(select(ClauseLibraryEntry)).all()}
    refs = [(r, 1) for r in find_clause_references("52.204-7 System for Award Management (OCT 2018) and 52.212-1 (SEP 2023) ALTERNATE I")]
    issues = {f"{m.reference.number}": m.issues for m in match_references(refs, library)}
    assert any(i.startswith("clause_date_differs") for i in issues["52.204-7"])
    assert any(i.startswith("clause_modified") for i in issues["52.212-1"])


# ── AC8: conflicting instructions surfaced ──


def test_conflicting_instructions_are_surfaced(session) -> None:
    o = _opp(session, set_aside_code=None, response_deadline=datetime(2026, 12, 1, 22, 0, tzinfo=UTC))
    f1 = _file(session, o, (QA / "solicitation.txt").read_text(encoding="utf-8"), "W912DY-26-R-0077_RFP.txt")
    f2 = _file(session, o, (QA / "questions_and_answers.txt").read_text(encoding="utf-8"), "W912DY-26-R-0077_Questions_and_Answers.txt")
    ids = {1: f1.id, 2: f2.id}
    provider = FakeProvider({"requirement_extraction_a": _remap(QA / "pass_a.json", ids), "requirement_extraction_b": _remap(QA / "pass_b.json", ids)})
    _run(session, o, provider, company_facts={})
    ten = _req(session, o, "not exceed 10 pages")
    fifteen = _req(session, o, "not exceed 15 pages")
    assert ten.status == "needs_review" and fifteen.status == "needs_review"
    assert "conflict_ambiguous" in ten.reconciliation["flags"]
    finding = session.scalar(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id, ComplianceFinding.finding_type == "conflict_ambiguous"))
    assert finding.blocks_submission and "page_limit" in finding.description
    injected = next(c for c in provider.calls if c["purpose"] == "requirement_extraction_a")
    assert "IGNORE ALL PREVIOUS INSTRUCTIONS" in injected["user"] and "IGNORE ALL PREVIOUS INSTRUCTIONS" not in injected["system"]
    assert "SOURCE SECURITY RULES" in injected["system"]
    assert not session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id, Requirement.status == "satisfied")).all()


# ── AC9: amendments invalidate stale conclusions ──


def test_amendment_invalidates_affected_conclusions(session) -> None:
    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids), supplier={"lead_time_days": 5, "transit_days": 3})
    delivery = _req(session, o, "within 30 calendar days")
    assert delivery.status == "satisfied"
    references = _req(session, o, "three past performance references")
    reviewer = _user(session, "reviewer")
    record_human_verification(session, requirement_id=references.id, actor=reviewer, description="Three references confirmed in the draft past performance volume.")
    _run(session, o, _dla_provider(ids), supplier={"lead_time_days": 5, "transit_days": 3})
    session.refresh(references)
    assert references.status == "satisfied"
    pursuit = Pursuit(opportunity_id=o.id, stage="drafting")
    session.add(pursuit)
    session.flush()
    proposal = Proposal(opportunity_id=o.id, pursuit_id=pursuit.id)
    session.add(proposal)
    session.flush()
    version = ProposalVersion(proposal_id=proposal.id, version_number=1, created_by="ai", full_text="x")
    session.add(version)
    session.flush()
    section = ProposalSection(proposal_version_id=version.id, section_key="past_performance", content="References A, B, C", requirement_ids=[references.id])
    session.add(section)
    session.flush()
    decisions_before = session.scalar(select(DecisionRun.id).where(DecisionRun.opportunity_id == o.id).order_by(DecisionRun.id.desc()).limit(1))

    ids[3] = _add_amendment(session, o).id
    result = _run(session, o, _dla_provider(ids), supplier={"lead_time_days": 5, "transit_days": 3})
    session.refresh(delivery)
    session.refresh(references)
    session.refresh(section)
    assert delivery.status == "superseded" and delivery.superseded_by_requirement_id is not None
    new_delivery = session.get(Requirement, delivery.superseded_by_requirement_id)
    assert "20 calendar days" in new_delivery.source_quote and new_delivery.amendment_changed_at is not None
    assert references.status == "stale" and references.stale_due_to_amendment and references.blocks_submission
    assert section.status == "stale"
    amendment = result["amendment"]
    assert references.id in amendment["previously_satisfied_ids"]
    assert amendment["impact"]["material"] and amendment["impact"]["rerun_jev_bid_decision"]
    alert = session.scalar(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id, ComplianceFinding.finding_type == "compliance_status_changed"))
    assert alert.description.startswith("COMPLIANCE STATUS CHANGED: 1 previously satisfied requirement requires revalidation because Amendment 0001")
    assert result["bid_decision_rerun_id"] is not None
    bundles = {r.bundle_name for r in session.scalars(select(DecisionRun).where(DecisionRun.opportunity_id == o.id, DecisionRun.id > decisions_before)).all()}
    assert {"compliance_and_amendment", "bid_decision"} <= bundles
    ack = _req(session, o, "acknowledge receipt of this amendment 0001")
    assert ack.requirement_type == "amendment_acknowledgment" and ack.severity == "critical"
    assert any(r["amendment_freshness"]["stale"] for r in compliance_matrix(session, o.id, filters=["stale"]))
    assert compliance_matrix(session, o.id, filters=["recently_changed_by_amendment"])


# ── AC10: critical requirements get redundant validation ──


def test_critical_requirements_need_two_validation_methods(session) -> None:
    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids))
    sam = _req(session, o, "System for Award Management (SAM) at the time")
    assert sam.severity == "critical"
    assert sam.status == "needs_review" and "redundant validation" in sam.status_reason
    registration = add_evidence(session, requirement_id=sam.id, evidence_type="company_registration", verification_method="company_record", verification_status="verified", description="SAM entity record: active through 2027-06-30")
    validator = {"validations": [{"requirement_id": sam.id, "status": "SATISFIED", "reason": "SAM record active", "evidence_refs": [{"evidence_id": registration.id}], "confidence": 0.93}]}
    provider = _dla_provider(ids, {"compliance_validator": validator})
    _run(session, o, provider)
    session.refresh(sam)
    assert sam.status == "satisfied"
    assert set(sam.validation["methods"]) == {"ai_validation", "deterministic"}

    single = ValidationInputs(requirement_type="certification", mandatory=True, severity="critical", has_source_location=True, current_status="unreviewed", stale=False, deterministic=[{"validator": "v", "status": "pass", "reason": "ok"}])
    assert decide_status(single).status == "needs_review"
    single.severity = "high"
    assert decide_status(single).status == "satisfied"


# ── AC11: coverage counts by category ──


def test_coverage_counts_are_visible_by_category(session) -> None:
    o, ids = _dla_setup(session)
    result = _run(session, o, _dla_provider(ids))
    counts = result["counts"]
    for key in ("mandatory_total", "mandatory_satisfied", "mandatory_missing", "mandatory_unknown", "mandatory_needs_review", "critical_total", "critical_satisfied", "critical_unresolved", "submission_blockers"):
        assert isinstance(counts[key], int)
    assert set(counts["by_category"]) == {"Administrative", "Technical", "Pricing", "Delivery", "Certifications", "Submission"}
    assert sum(v["total"] for v in counts["by_category"].values()) == counts["mandatory_total"]
    assert counts["percentages"]["definition"]
    run = session.get(ComplianceRun, result["matrix_run_id"])
    assert run.mandatory_total == counts["mandatory_total"] and run.critical_unresolved == counts["critical_unresolved"]
    assert run.false_satisfied_detected == 0

    from govcon.cli import app

    session.commit()
    output = CliRunner().invoke(app, ["compliance", "coverage", "--opportunity-id", str(o.id)])
    assert output.exit_code == 0, output.output
    assert "Mandatory requirements:" in output.output and "Certifications" in output.output


# ── AC12: satisfied requires evidence or validator output ──


def test_satisfied_always_has_evidence_or_validator_output(session) -> None:
    req = Requirement(opportunity_id=_opp(session).id, requirement_text="x" * 10, status="unreviewed")
    with pytest.raises(ComplianceInvariantError):
        apply_decision(req, StatusDecision("satisfied", "no basis", False, [], False))
    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids), supplier={"lead_time_days": 5, "transit_days": 3})
    satisfied = session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id, Requirement.status == "satisfied")).all()
    assert satisfied
    for r in satisfied:
        assert r.validation["methods"], r.requirement_text
        row = next(x for x in compliance_matrix(session, o.id) if x["requirement_id"] == r.id)
        assert row["evidence"] or any(v["status"] == "pass" for v in row["validator_output"])


# ── AC13: red-team findings persisted ──


def test_red_team_findings_are_persisted(session) -> None:
    o, ids = _dla_setup(session)
    red_team = {"findings": [
        {"finding_type": "unsigned_form", "certainty": "confirmed", "severity": "critical", "description": "SF 1449 signature is not evidenced.", "evidence": [{"source_file_id": ids[1], "quote": "Quoters shall complete and sign the SF 1449"}]},
        {"finding_type": "delivery_mismatch", "certainty": "possible", "severity": "medium", "description": "Transit time not documented.", "missing_evidence": "carrier transit estimate"},
    ]}
    provider = _dla_provider(ids, {"compliance_red_team": red_team})
    result = _run(session, o, provider)
    rows = session.scalars(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id, ComplianceFinding.compliance_run_id == result["red_team"]["run_id"])).all()
    ai_rows = [r for r in rows if r.detected_by == "compliance_red_team_ai"]
    assert {(r.finding_type, r.certainty, r.blocks_submission) for r in ai_rows} == {("unsigned_form", "confirmed", True), ("delivery_mismatch", "possible", False)}
    assert any(r.detected_by == "red_team_rules" and r.finding_type == "critical_requirement_unresolved" for r in rows)
    red = next(c for c in provider.calls if c["purpose"] == "compliance_red_team")
    extract = next(c for c in provider.calls if c["purpose"] == "requirement_extraction_a")
    assert red["system"] != extract["system"] and "adversarial government-bid compliance reviewer" in red["system"]


# ── AC14: proposal coverage checked automatically ──


def _proposal(session, opp, sections: list[tuple[str, str, list[int]]]) -> ProposalVersion:
    pursuit = session.scalar(select(Pursuit).where(Pursuit.opportunity_id == opp.id))
    if pursuit is None:
        pursuit = Pursuit(opportunity_id=opp.id, stage="drafting")
        session.add(pursuit)
        session.flush()
    proposal = session.scalar(select(Proposal).where(Proposal.opportunity_id == opp.id))
    if proposal is None:
        proposal = Proposal(opportunity_id=opp.id, pursuit_id=pursuit.id)
        session.add(proposal)
        session.flush()
    number = 1 + len(session.scalars(select(ProposalVersion).where(ProposalVersion.proposal_id == proposal.id)).all())
    version = ProposalVersion(proposal_id=proposal.id, version_number=number, created_by="ai", full_text="\n".join(c for _, c, _ in sections))
    session.add(version)
    session.flush()
    for order, (key, content, req_ids) in enumerate(sections):
        session.add(ProposalSection(proposal_version_id=version.id, section_key=key, heading=key, sort_order=order, content=content, requirement_ids=req_ids))
    session.flush()
    return version


def test_proposal_coverage_is_checked_against_the_matrix(session) -> None:
    from govcon.compliance.proposal_coverage import check_proposal_coverage
    from govcon.compliance.validator import run_validation

    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids))
    refs = _req(session, o, "three past performance references")
    delivery = _req(session, o, "within 30 calendar days")
    two = _proposal(session, o, [
        ("past_performance", "Past performance references for similar supply contracts:\nReference A: DLA glove contract\nReference B: VA supply contract", [refs.id]),
        ("delivery", "We will provide excellent customer service.", [delivery.id]),
    ])
    result = check_proposal_coverage(session, o.id, two.id)
    assert result["results"][str(refs.id)]["coverage_status"] == "PARTIAL"
    assert result["results"][str(refs.id)]["detected_count"] == 2
    assert result["results"][str(delivery.id)]["coverage_status"] == "NEEDS_REVIEW", "writer mapping without substance must not count"
    gap = session.scalar(select(ComplianceFinding).where(ComplianceFinding.requirement_id == refs.id, ComplianceFinding.finding_type == "proposal_coverage_gap", ComplianceFinding.status == "open"))
    assert gap is not None and gap.blocks_submission and gap.description.startswith("BLOCKING")
    run_validation(session, o.id, use_ai=False)
    session.refresh(refs)
    assert refs.status == "missing"

    three = _proposal(session, o, [("past_performance", "Past performance references for similar supply contracts:\nReference A: DLA\nReference B: VA\nReference C: USACE", [refs.id])])
    result = check_proposal_coverage(session, o.id, three.id)
    assert result["results"][str(refs.id)]["coverage_status"] == "COVERED"
    session.refresh(gap)
    assert gap.status == "resolved"


# ── AC15 / AC16 / AC17: pre-flight, readiness blocking, audited overrides ──


def test_preflight_blocks_ready_to_submit_and_override_is_audited(session, tmp_path) -> None:
    from govcon.compliance.submission_preflight import ReadinessBlocked, move_to_ready_to_submit, run_submission_preflight

    o, ids = _dla_setup(session, with_amendment=True)
    _run(session, o, _dla_provider(ids))
    pursuit = Pursuit(opportunity_id=o.id, stage="review", approved_to_bid_at=datetime.now(UTC))  # bid approved
    session.add(pursuit)
    session.flush()
    submission = Submission(opportunity_id=o.id, pursuit_id=pursuit.id)
    session.add(submission)
    session.flush()
    quote = tmp_path / "quote.docx"
    from govcon.proposals.service import get_or_create_proposal
    from govcon.proposals.versions import create_proposal_version
    from govcon.proposals.export import export_proposal_docx
    proposal = get_or_create_proposal(session, opportunity_id=o.id, pursuit_id=pursuit.id)
    version = create_proposal_version(session, proposal_id=proposal.id, created_by="synthetic fixture",
                                     status_after="ai_generated", sections=[{"section_key":"quote", "content":"Fixture quote content"}])
    content = export_proposal_docx(session, proposal_version_id=version.id)
    quote.write_bytes(content)
    package = SubmissionPackage(files=[PackageFile(
        "quote.docx", role="proposal", size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(), local_path=str(quote),
    )], proposal_version_id=version.id, recipient_email="jane.buyer@dla.mil", amendments_acknowledged=[])
    result = run_submission_preflight(session, o.id, package, now=NOW)
    items = {i["check"]: i for i in result["items"]}
    assert result["ready"] is False
    for check in ("proposal_content", "pricing_workbook", "signed_documents", "representations", "certifications", "amendment_acknowledgments", "required_attachments", "file_types", "page_limit", "recipient", "deadline", "timezone", "submission_instructions", "matrix_blockers"):
        assert check in items, check
    assert items["amendment_acknowledgments"]["status"] == "fail"
    assert items["file_types"]["status"] == "fail"
    assert items["recipient"]["status"] == "pass"
    assert items["timezone"]["status"] == "pass"
    assert items["page_limit"]["status"] == "unknown"

    approver = _user(session, "approver")
    reviewer = _user(session, "reviewer")
    with pytest.raises(ReadinessBlocked) as blocked:
        move_to_ready_to_submit(session, o.id, actor=approver)
    assert any(b["kind"] == "requirement" for b in blocked.value.blockers)
    from govcon.collaboration.users import PermissionDenied

    with pytest.raises(PermissionDenied):
        move_to_ready_to_submit(session, o.id, actor=reviewer, override_reason="Reviewer wants to push it through.")
    with pytest.raises(ReadinessBlocked):
        move_to_ready_to_submit(session, o.id, actor=approver, override_reason="short")
    session.refresh(pursuit)
    assert pursuit.stage == "review"

    moved = move_to_ready_to_submit(session, o.id, actor=approver, override_reason="Owner accepts residual risk after CO phone confirmation.")
    assert moved.stage == "ready_to_submit"
    audit = session.scalar(select(AuditEvent).where(AuditEvent.opportunity_id == o.id, AuditEvent.action_type == "compliance_readiness_override"))
    assert audit.user_id == approver.id and audit.new_value["reason"].startswith("Owner accepts") and audit.old_value["blockers"]


def test_override_requirement_needs_role_reason_version_and_audit(session) -> None:
    from govcon.collaboration.users import PermissionDenied
    from govcon.concurrency import StaleRecordError

    o, ids = _dla_setup(session)
    _run(session, o, _dla_provider(ids))
    req = _req(session, o, "small business under NAICS 339113")
    reviewer, owner = _user(session, "reviewer"), _user(session, "owner")
    with pytest.raises(PermissionDenied):
        override_requirement(session, requirement_id=req.id, status="satisfied", actor=reviewer, reason="Size certified in SAM today.", expected_version=req.version)
    with pytest.raises(ValueError):
        override_requirement(session, requirement_id=req.id, status="satisfied", actor=owner, reason="ok", expected_version=req.version)
    with pytest.raises(StaleRecordError):
        override_requirement(session, requirement_id=req.id, status="satisfied", actor=owner, reason="Size certified in SAM today.", expected_version=req.version - 1)
    updated = override_requirement(session, requirement_id=req.id, status="satisfied", actor=owner, reason="Size certified in SAM today.", expected_version=req.version)
    assert updated.status == "satisfied" and updated.verified_by_human
    audit = session.scalar(select(AuditEvent).where(AuditEvent.entity_type == "requirements", AuditEvent.entity_id == req.id, AuditEvent.action_type == "compliance_requirement_override"))
    assert audit.old_value["status"] == "unknown" and audit.new_value["reason"] == "Size certified in SAM today."
    _run(session, o, _dla_provider(ids))
    session.refresh(updated)
    assert updated.status == "satisfied", "later validation keeps an authorized override"


def test_ready_to_submit_succeeds_only_when_everything_is_green(session, tmp_path) -> None:
    from govcon.compliance.proposal_coverage import check_proposal_coverage
    from govcon.compliance.submission_preflight import move_to_ready_to_submit, readiness_blockers, run_submission_preflight
    from govcon.compliance.validator import run_validation

    o = _opp(session, set_aside_code=None, response_deadline=datetime(2026, 10, 20, 21, 0, tzinfo=UTC))
    text = (
        "SOLICITATION RFQ 55\nSECTION 1 INSTRUCTIONS\n"
        "Quotes are due no later than October 20, 2026 5:00 PM ET.\n"
        "Quotes shall be submitted by email to buyer@agency.gov.\n"
        "The quoter shall describe the proposed nitrile glove product including manufacturer and model.\n"
    )
    _file(session, o, text, "RFQ_55_solicitation.txt")
    _run(session, o, FakeProvider(), company_facts={}, use_ai=False)
    owner = _user(session, "owner")
    pursuit = Pursuit(opportunity_id=o.id, stage="review", approved_to_bid_at=datetime.now(UTC))  # bid approved
    session.add(pursuit)
    session.flush()
    product = _req(session, o, "proposed nitrile glove product")
    version = _proposal(session, o, [("technical", "Proposed nitrile glove product: manufacturer Acme, model N-100, described in full.", [product.id])])
    check_proposal_coverage(session, o.id, version.id)
    run_validation(session, o.id, use_ai=False)
    for req in session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id, Requirement.status.not_in(["satisfied", "superseded", "not_applicable"]))).all():
        override_requirement(session, requirement_id=req.id, status="satisfied", actor=owner, reason="Verified by owner against the RFQ text.", expected_version=req.version, acknowledge_deterministic_failure=True)
    for finding in session.scalars(select(ComplianceFinding).where(ComplianceFinding.opportunity_id == o.id, ComplianceFinding.status == "open", ComplianceFinding.blocks_submission.is_(True))).all():
        from govcon.compliance.matrix import resolve_finding

        resolve_finding(session, finding.id, actor=owner, notes="Reviewed and accepted by owner.")
    submission = Submission(opportunity_id=o.id, pursuit_id=pursuit.id, recipient_email="buyer@agency.gov", deadline_timezone="ET", required_files=[])
    session.add(submission)
    session.flush()
    quote = tmp_path / "quote.pdf"
    from govcon.proposals.export import export_proposal_pdf

    # The proposal file must be bytes exported from the pinned version (F26).
    quote.write_bytes(export_proposal_pdf(session, proposal_version_id=version.id))
    package = SubmissionPackage(files=[PackageFile("quote.pdf", role="proposal", size_bytes=quote.stat().st_size, page_count=2, sha256=hashlib.sha256(quote.read_bytes()).hexdigest(), local_path=str(quote))], recipient_email="buyer@agency.gov", amendments_acknowledged=[], proposal_version_id=version.id)
    result = run_submission_preflight(session, o.id, package, submission_id=submission.id, now=NOW)
    assert result["deterministic_ready"], [i for i in result["items"] if i["status"] not in {"pass", "not_applicable"}]
    assert result["status"] == "ready", (result["status"], result["jev"], result["ai"])
    assert readiness_blockers(session, o.id) == []
    assert move_to_ready_to_submit(session, o.id, actor=owner).stage == "ready_to_submit"
    assert submission.readiness_status == "ready"
    assert not session.scalar(select(AuditEvent).where(AuditEvent.opportunity_id == o.id, AuditEvent.action_type == "compliance_readiness_override"))

    requirement = session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id)).first()
    requirement.version += 1
    session.flush()
    assert any("changed after the latest pre-flight" in b["description"] for b in readiness_blockers(session, o.id))


# ── AC18 / AC19 / AC20: benchmark, metrics, release gate ──


def test_compliance_benchmark_passes_and_metrics_are_measurable() -> None:
    from govcon.compliance.regression import run_benchmark_suite

    suite = run_benchmark_suite(FIXTURES)
    assert suite.gate.passed, suite.gate.failures
    assert {c.case_id for c in suite.cases} == {"dla_base_amendment", "table_embedded_packaging", "qa_conflict_injection"}
    for metric in ("mandatory_recall", "critical_recall", "citation_accuracy", "false_satisfied_rate", "amendment_change_detection"):
        assert suite.aggregate[metric] is not None
    assert suite.aggregate["critical_recall"] == 1.0 and suite.aggregate["false_satisfied_rate"] == 0.0
    assert 0.9 < suite.aggregate["citation_accuracy"] < 1.0, "one fixture citation is deliberately inaccurate"
    table = next(c for c in suite.cases if c.case_id == "table_embedded_packaging")
    assert table.details["found_by"]["packaging_mil_std"] == ["D"], "table-embedded requirement caught only by the deterministic scan"
    for case_dir in (DLA, TABLE, QA):
        case = json.loads((case_dir / "case.json").read_text(encoding="utf-8"))
        for key in ("source_files", "expected_mandatory_requirements", "expected_critical_requirements", "expected_conflicts", "expected_amendment_changes", "expected_submission_files"):
            assert key in case


def test_critical_recall_regression_blocks_release(tmp_path) -> None:
    from govcon.cli import app
    from govcon.compliance.regression import run_benchmark_suite

    root = tmp_path / "compliance"
    shutil.copytree(FIXTURES, root)
    with patch("govcon.compliance.regression.scan_requirements", return_value=[]):
        suite = run_benchmark_suite(root)
    assert not suite.gate.passed
    assert any("critical requirement recall regression" in f and "packaging_mil_std" in f for f in suite.gate.failures)

    case = json.loads((root / "qa_conflict_injection" / "case.json").read_text(encoding="utf-8"))
    case["expected_mandatory_requirements"].append({"id": "offer_acceptance_period", "match": ["acceptance period of 90 days"], "critical": True})
    (root / "qa_conflict_injection" / "case.json").write_text(json.dumps(case), encoding="utf-8")
    result = CliRunner().invoke(app, ["compliance", "benchmark", "--fixtures", str(root)])
    assert result.exit_code == 1 and "FAIL" in result.output and "offer_acceptance_period" in result.output


# ── prompts, structured runner, extraction fix ──


def test_compliance_prompts_are_registry_managed_and_gated(session) -> None:
    from govcon.prompting.evaluation import run_activation_gate
    from govcon.prompting.loader import load_markdown_prompt
    from govcon.prompting.registry import PromptActivationBlocked, activate_prompt, active_version, rollback_prompt, sync_prompts
    from govcon.config import get_settings

    synced = sync_prompts(session, PROMPT_ROOT)
    for name in COMPLIANCE_PROMPTS:
        assert name in synced and active_version(session, name) == "v1"
        path = next(PROMPT_ROOT.rglob(f"{name}_v1.md"))
        gate = run_activation_gate(load_markdown_prompt(path), PROMPT_ROOT, settings=get_settings(), run_regression=name == "requirement_extraction_a")
        assert gate.passed, gate.failures

    candidate = PROMPT_ROOT / "deepseek" / "requirement_extraction_a_v2.md"
    try:
        candidate.write_text((PROMPT_ROOT / "deepseek" / "requirement_extraction_a_v1.md").read_text(encoding="utf-8").replace("version: v1", "version: v2").replace("status: active", "status: candidate"))
        sync_prompts(session, PROMPT_ROOT)
        result = activate_prompt(session, "requirement_extraction_a", "v2", prompt_root=PROMPT_ROOT)
        assert result["gate"]["required"] and result["gate"]["passed"] and result["previous_version"] == "v1"
        assert active_version(session, "requirement_extraction_a") == "v2"
        assert rollback_prompt(session, "requirement_extraction_a")["active_version"] == "v1"

        candidate.write_text(candidate.read_text(encoding="utf-8").replace("required_variables: OPPORTUNITY_JSON, DOCUMENT_INVENTORY_JSON, SOURCE_CHUNKS, AMENDMENT_JSON\n", "").replace("version: v2", "version: v3"))
        candidate.rename(candidate.with_name("requirement_extraction_a_v3.md"))
        sync_prompts(session, PROMPT_ROOT)
        with pytest.raises(PromptActivationBlocked):
            activate_prompt(session, "requirement_extraction_a", "v3", prompt_root=PROMPT_ROOT)
        assert active_version(session, "requirement_extraction_a") == "v1"
    finally:
        for leftover in ("requirement_extraction_a_v2.md", "requirement_extraction_a_v3.md"):
            (PROMPT_ROOT / "deepseek" / leftover).unlink(missing_ok=True)


def test_no_hardcoded_prompts_in_compliance_services() -> None:
    root = Path(__file__).parent.parent / "src" / "govcon" / "compliance"
    for py_file in root.glob("*.py"):
        for node in ast.walk(ast.parse(py_file.read_text(encoding="utf-8"))):
            if isinstance(node, ast.Constant) and isinstance(node.value, str) and len(node.value) > 200:
                assert "you are" not in node.value.lower(), f"hardcoded prompt in {py_file.name}"


def test_structured_runner_fails_closed_on_malformed_output(session) -> None:
    from govcon.ai.structured import StructuredCallError, run_structured_prompt
    from govcon.security.classification import DataClassification

    o = _opp(session)
    for bad in ("not json", json.dumps({"requirements": [{"requirement_text": "x"}]}), json.dumps({"validations": [{"requirement_id": 1, "status": "SATISFIED", "reason": "no evidence"}]})):
        prompt = "compliance_validator" if "validations" in bad else "requirement_extraction_b"
        variables = {"REQUIREMENTS_JSON": [], "EVIDENCE_JSON": []} if prompt == "compliance_validator" else {"DOCUMENT_INVENTORY_JSON": [], "SOURCE_CHUNKS": "x", "AMENDMENT_JSON": {}}
        with patch("govcon.ai.structured.get_provider", return_value=FakeProvider({prompt: bad})):
            with pytest.raises(StructuredCallError):
                run_structured_prompt(session, opportunity_id=o.id, prompt_name=prompt, analysis_type="compliance_review", variables=variables, context_manifest={}, classification=DataClassification.PUBLIC)
    assert not session.scalars(select(AIAnalysis).where(AIAnalysis.opportunity_id == o.id)).all()
    with patch("govcon.ai.structured.get_provider", return_value=FakeProvider()):
        with pytest.raises(StructuredCallError) as exc:
            run_structured_prompt(session, opportunity_id=o.id, prompt_name="requirement_extraction_b", analysis_type="compliance_review", variables={"SOURCE_CHUNKS": "x"}, context_manifest={}, classification=DataClassification.PUBLIC)
    assert exc.value.reason == "render_error"


def test_no_ai_key_run_is_incomplete_not_green(session) -> None:
    o, _ = _dla_setup(session)
    from govcon.ai.providers import NoProviderConfigured

    with patch("govcon.ai.structured.get_provider", side_effect=NoProviderConfigured("DEEPSEEK_API_KEY is required")):
        result = run_compliance_pipeline(session, o.id, company_facts={}, now=NOW)
    assert result["status"] == "incomplete"
    assert any(w["code"].startswith("pass_a_no_provider") for w in result["warnings"])
    reqs = session.scalars(select(Requirement).where(Requirement.opportunity_id == o.id)).all()
    assert reqs and not any(r.independently_confirmed for r in reqs)
    assert not any(r.status == "satisfied" and r.severity == "critical" for r in reqs)


def test_docx_tables_are_extracted() -> None:
    import io

    from docx import Document

    from govcon.enrich.extract import extract_text

    doc = Document()
    doc.add_paragraph("ATTACHMENT 7 PACKAGING")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Requirement"
    table.cell(1, 0).text, table.cell(1, 1).text = "Packaging", "Items shall be packaged per MIL-STD-2073-1, Level A."
    buffer = io.BytesIO()
    doc.save(buffer)
    result = extract_text(buffer.getvalue(), "", "attachment_7.docx")
    assert result.status == "success"
    assert "[Table 1]" in result.text and "MIL-STD-2073-1, Level A" in result.text
